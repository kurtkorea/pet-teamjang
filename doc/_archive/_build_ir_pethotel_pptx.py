# -*- coding: utf-8 -*-
"""PAWKOREA Pet Hotel Yeongjong IR Deck v0.4 — thesis-first narrative, editorial layout, facility facts from the original business plan."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

PAPER = RGBColor(0xF5, 0xF3, 0xEE)
INK = RGBColor(0x14, 0x14, 0x14)
GRAY = RGBColor(0x6B, 0x6B, 0x6B)
RULE = RGBColor(0xD9, 0xD5, 0xCC)
SAND = RGBColor(0xEA, 0xE6, 0xDD)
ACCENT = RGBColor(0xFF, 0x5B, 0x1F)
ACCENT_DIM = RGBColor(0xFF, 0xC9, 0xB3)
ACCENT_WASH = RGBColor(0xFF, 0xEC, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NIGHT = RGBColor(0x14, 0x14, 0x14)
NIGHT_GRAY = RGBColor(0x9A, 0x9A, 0x9A)

W, H = Inches(13.333), Inches(7.5)
MX = 0.7
FONT = "Malgun Gothic"
OUT = r"D:\dev_pawkorea\pet-teamjang\doc\IR_PAWKOREA_pet_hotel_yeongjong_v0.4.pptx"

# 객실 구성: 원 사업계획서(리조이스) 2F·3F 평면 기준
ROOMS = [("소형", 52, 50_000), ("중형", 34, 65_000), ("대형", 22, 90_000), ("패밀리", 4, 150_000)]
ADDON = 0.15
OCC_SCENARIOS = (0.30, 0.45, 0.60)


def I(v):
    return Inches(v)


def _font(run, size, color, bold):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, color, bold) or (text, size, color, bold, space_after)."""
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        t, size, color, bold = line[:4]
        after = line[4] if len(line) > 4 else 6
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        r = p.add_run()
        r.text = t
        _font(r, size, color, bold)
    return tb


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def hline(slide, x, y, w, color=RULE, weight=0.015):
    return box(slide, x, y, w, weight, color)


def vline(slide, x, y, h, color=RULE, weight=0.015):
    return box(slide, x, y, weight, h, color)


def dot(slide, cx, cy, d, color):
    return box(slide, cx - d / 2, cy - d / 2, d, d, color, MSO_SHAPE.OVAL)


def new_slide(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = NIGHT if dark else PAPER
    return s


def frame(slide, page, section, title, dark=False):
    """Section label, action title, page footer."""
    fg = WHITE if dark else INK
    sub = NIGHT_GRAY if dark else GRAY
    text(slide, MX, 0.45, 6, 0.3, [(section, 11, ACCENT, True)])
    text(slide, MX, 0.8, 11.9, 1.0, [(title, 25, fg, True)])
    text(slide, MX, 7.0, 6, 0.25, [("PAWKOREA Pet Hotel Yeongjong", 9, sub, False)])
    text(slide, 11.63, 7.0, 1.0, 0.25, [(f"{page:02d}", 9, sub, True)], PP_ALIGN.RIGHT)


def style_chart(chart, number_format, colors):
    chart.has_legend = False
    chart.has_title = False
    chart.font.size = Pt(12)
    chart.font.name = FONT
    plot = chart.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(14)
    labels.font.bold = True
    labels.font.color.rgb = INK
    series = plot.series[0]
    for i, color in enumerate(colors):
        pt = series.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = color
    va = chart.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.minimum_scale = 0
    ca = chart.category_axis
    ca.tick_labels.font.size = Pt(12)
    ca.tick_labels.font.color.rgb = GRAY
    ca.format.line.color.rgb = RULE
    ca.has_major_gridlines = False


def room_economics():
    total = sum(n for _, n, _ in ROOMS)
    adr = sum(n * p for _, n, p in ROOMS) / total
    full_year = total * 365 * adr * (1 + ADDON)
    revenues = [full_year * occ / 1e8 for occ in OCC_SCENARIOS]
    return total, adr, revenues


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    total_rooms, adr, revenues = room_economics()
    step_10pt = (revenues[1] - revenues[0]) / ((OCC_SCENARIOS[1] - OCC_SCENARIOS[0]) * 10)
    page = 0

    # 01 COVER
    page += 1
    s = new_slide(prs)
    box(s, 0, 0, 13.333, 0.12, ACCENT)
    text(s, MX, 0.7, 8, 0.3, [("INVESTOR BRIEF  ·  2026.10  ·  v0.4  ·  논의용 초안", 11, GRAY, True)])
    text(
        s,
        MX,
        1.6,
        11.5,
        2.4,
        [
            ("맡기는 순간부터,", 54, INK, True, 0),
            ("보이는 안심을.", 54, ACCENT, True, 0),
        ],
    )
    text(
        s,
        MX,
        3.95,
        10,
        0.9,
        [
            ("파우코리아 펫호텔 영종", 22, INK, True, 4),
            ("공사를 마친 영종 위탁형 펫호텔과 자체 AI 관제 PET-Aeye를 함께 운영합니다.", 15, GRAY, False),
        ],
    )
    hline(s, MX, 5.55, 11.93, INK, 0.02)
    facts = [
        (f"{total_rooms}실", "4개 타입 위탁 객실"),
        ("15분", "인천공항까지"),
        ("공사 완료", "1차 1동 4층 · 약 370평"),
        ("PET-Aeye", "실시간 AI 관제 솔루션"),
    ]
    for i, (v, lab) in enumerate(facts):
        x = MX + i * 3.0
        text(s, x, 5.8, 2.8, 0.5, [(v, 24, INK, True, 2), (lab, 12, GRAY, False)])

    # 02 THESIS
    page += 1
    s = new_slide(prs)
    frame(s, page, "INVESTMENT THESIS", "이 사업을 제안하는 세 가지 이유")
    thesis = [
        (
            "01",
            "영종에는 구조적인 위탁 수요가 있습니다",
            "인스파이어 리조트는 반려동물 입실이 어렵고, 공항 15분 거리와 신도시 생활권이 한 섬에 겹쳐 있습니다.",
        ),
        (
            "02",
            "공식 공항 채널 밖에 중대형 수요가 남아 있습니다",
            "인천공항 공식 호텔링은 1개 업체가 운영하며 약 13kg 이하 소형견 중심입니다.",
        ),
        (
            "03",
            "객실이 아니라 '보이는 안심'으로 차별화합니다",
            "자체 AI 관제 PET-Aeye가 실시간 캠 · AI 행동 이벤트 · 일일 리포트를 전 객실에 제공합니다.",
        ),
    ]
    for i, (n, head, proof) in enumerate(thesis):
        y = 2.05 + i * 1.55
        hline(s, MX, y - 0.2, 11.93)
        text(s, MX, y, 1.0, 0.8, [(n, 34, ACCENT, True)])
        text(s, MX + 1.3, y + 0.05, 10.5, 1.2, [(head, 19, INK, True, 6), (proof, 13, GRAY, False)])
    hline(s, MX, 2.05 + 3 * 1.55 - 0.2, 11.93)

    # 03 WEDGE
    page += 1
    s = new_slide(prs)
    frame(s, page, "THE WEDGE", "세 가지 숫자에서 진입 기회를 찾았습니다")
    wedge = [
        ("0실", "인스파이어 리조트의\n반려동물 입실 가능 객실", "보조견 외에는 입실이 어려워, 투숙객에게 인근 위탁이 필요합니다."),
        ("13kg", "인천공항 공식 호텔링의\n수용 체중 기준(안내)", "중·대형견 보호자는 다른 선택지를 찾아야 합니다."),
        ("26실", "저희 3층의\n대형 22실 · 패밀리 4실", "공식 채널이 받기 어려운 보호자를 맞이합니다."),
    ]
    for i, (num, lab, note) in enumerate(wedge):
        x = MX + i * 4.1
        if i:
            vline(s, x - 0.3, 2.1, 4.2)
        text(s, x, 2.0, 3.8, 1.4, [(num, 72, ACCENT if i == 2 else INK, True)])
        text(s, x, 3.6, 3.7, 1.0, [(line, 16, INK, True, 2) for line in lab.split("\n")])
        text(s, x, 4.75, 3.6, 1.2, [(note, 13, GRAY, False)])
    text(
        s,
        MX,
        6.45,
        11.9,
        0.4,
        [("Source: 인스파이어 이용약관·FAQ, 인천국제공항공사 보도자료(2024.3·11) 및 영문 안내, 원 사업계획서 3F 평면", 10, GRAY, False)],
    )

    # 04 DEMAND & MARKET
    page += 1
    s = new_slide(prs)
    frame(s, page, "DEMAND", "시장은 아직 작지만 빠르게 크고, 입지가 수요를 만듭니다")
    streams = [
        ("공항 출국", "출국 전 맡기고 귀국 시 찾는 수요입니다. 공항공사가 호텔링 카운터를 직접 열 만큼 수요가 확인됐습니다."),
        ("리조트 투숙", "반려동물 입실이 어려운 인스파이어 투숙객의 단기 위탁 수요입니다."),
        ("영종 생활권", "하늘도시·미단시티 거주 가구의 출장·여행·데이케어 수요로 비수기를 보완합니다."),
    ]
    for i, (h, b) in enumerate(streams):
        y = 2.0 + i * 1.45
        dot(s, MX + 0.12, y + 0.17, 0.18, ACCENT)
        text(s, MX + 0.45, y, 5.3, 1.3, [(h, 17, INK, True, 4), (b, 13, GRAY, False)])
    data = CategoryChartData()
    data.categories = ["2024", "2030E"]
    data.add_series("국내 펫 보딩", (980, 1830))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, I(7.1), I(2.0), I(5.4), I(3.7), data)
    style_chart(gf.chart, '#,##0"억 원"', [SAND, ACCENT])
    text(
        s,
        7.1,
        5.8,
        5.5,
        0.9,
        [
            ("국내 펫 보딩 시장 · CAGR 11.1% (2025–2030 전망)", 13, INK, True, 2),
            ("Grand View Horizon, USD→KRW 1,350원 환산(약수). 동반숙박 시장은 제외했습니다.", 10, GRAY, False),
        ],
    )

    # 05 CUSTOMER JOURNEY
    page += 1
    s = new_slide(prs)
    frame(s, page, "FLAGSHIP PRODUCT", "‘공항 안심 보딩’ — 보호자의 여행 하루를 돕습니다")
    steps = [
        ("D-7", "예약", "룸 타입·체중 선택\n안심 옵션 추가"),
        ("D0", "체크인", "1층 프런트·상담실\n주차 후 입실"),
        ("D0", "샌딩", "공항까지 약 15분\n자체 차량(검토)"),
        ("여행 중", "보이는 안심", "배정 룸 라이브\n매일 AI 리포트"),
        ("귀국", "픽업", "공항 또는 시설\n수령"),
        ("D+1", "하이라이트", "투숙 요약 클립\n재방문 안내"),
    ]
    y = 3.0
    hline(s, MX + 0.4, y, 11.1, INK, 0.02)
    gap = 11.1 / (len(steps) - 1)
    for i, (when, what, desc) in enumerate(steps):
        cx = MX + 0.4 + i * gap
        hot = i == 3
        dot(s, cx, y + 0.01, 0.42 if hot else 0.26, ACCENT if hot else INK)
        text(s, cx - 1.0, 2.15, 2.0, 0.4, [(when, 12, ACCENT if hot else GRAY, True)], PP_ALIGN.CENTER)
        text(s, cx - 1.0, 3.45, 2.0, 0.45, [(what, 16, INK, True)], PP_ALIGN.CENTER)
        text(s, cx - 1.0, 3.95, 2.0, 1.0, [(d, 12, GRAY, False, 2) for d in desc.split("\n")], PP_ALIGN.CENTER)
    box(s, MX, 5.35, 11.93, 1.2, SAND)
    text(
        s,
        MX + 0.35,
        5.55,
        11.3,
        0.9,
        [
            ("공식 공항 카운터와 경쟁하지 않고, 보완하는 선택지가 됩니다.", 15, INK, True, 4),
            ("자체 샌딩 · 중대형 수용 · PET-Aeye 안심 리포트로, 공식 채널이 담기 어려운 보호자를 돕습니다.", 13, GRAY, False),
        ],
    )

    # 06 STATEMENT (dark)
    page += 1
    s = new_slide(prs, dark=True)
    text(s, MX, 0.45, 6, 0.3, [("OUR APPROACH", 11, ACCENT, True)])
    text(
        s,
        MX,
        2.0,
        11.5,
        2.4,
        [
            ("객실은 다른 곳에도 있습니다.", 42, WHITE, True, 6),
            ("저희는 '지금 괜찮다'는 확신을 함께 드립니다.", 30, ACCENT, True),
        ],
    )
    text(
        s,
        MX,
        4.8,
        11.3,
        1.6,
        [
            ("KB금융지주 경영연구소 조사에서 펫테크 기기를 쓰는 첫 번째 이유는 '혼자 두고 외출할 때의 걱정 감소'였고,", 14, NIGHT_GRAY, False, 2),
            ("모니터링 CCTV·카메라는 사용 기기 중 약 23%를 차지했습니다. PET-Aeye는 이 걱정을 숙박 중에도 덜어드립니다.", 14, NIGHT_GRAY, False),
        ],
    )
    text(s, 11.63, 7.0, 1.0, 0.25, [(f"{page:02d}", 9, NIGHT_GRAY, True)], PP_ALIGN.RIGHT)

    # 07 PRODUCT — PETI
    page += 1
    s = new_slide(prs)
    frame(s, page, "TECHNOLOGY", "PET-Aeye — 보고, 알리고, 요약하는 자체 AI 관제")
    loop = [
        ("보고", "룸별 카메라 라이브\n검출 오버레이"),
        ("알리고", "장시간 무움직임·이상 행동\n스태프 우선 알림"),
        ("요약", "Stay 단위 일일 리포트\n보호자 기간 한정 링크"),
    ]
    for i, (h, b) in enumerate(loop):
        x = MX + i * 2.05
        box(s, x, 2.0, 1.85, 1.85, ACCENT if i == 0 else INK, MSO_SHAPE.OVAL)
        text(s, x, 2.0, 1.85, 1.85, [(h, 20, WHITE, True)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(s, x - 0.1, 4.05, 2.05, 1.0, [(t, 11, GRAY, False, 2) for t in b.split("\n")], PP_ALIGN.CENTER)
    text(
        s,
        MX,
        5.05,
        6.1,
        1.0,
        [("원 사업계획의 '24시간 CHECKING SYSTEM'을 사람의 순찰과 AI 관제가 함께 맡습니다.", 12, INK, False)],
    )
    vline(s, 7.0, 1.95, 4.5)
    text(s, 7.35, 1.95, 5.3, 0.35, [("스태프용 관제", 13, ACCENT, True)])
    staff = [
        f"{total_rooms}실 룸 그리드 관제 · 이벤트 타임라인",
        "AI 검출 · 추적 · 행동 라벨(서기·걷기·눕기 등)",
        "이상 행동 우선 알림 · PTZ 자동 추적",
    ]
    text(s, 7.35, 2.35, 5.3, 1.6, [(f"✓  {t}", 13, INK, False, 5) for t in staff])
    text(s, 7.35, 3.95, 5.3, 0.35, [("보호자용 안심", 13, ACCENT, True)])
    guardian = [
        "기간 한정 링크로 배정 룸 라이브 열람",
        "Stay 단위 일일 CareReport",
        "퇴실 후 하이라이트 클립",
    ]
    text(s, 7.35, 4.35, 5.3, 1.6, [(f"✓  {t}", 13, INK, False, 5) for t in guardian])
    text(s, MX, 6.45, 11.9, 0.35, [("구성: RTSP 수신·녹화 · YOLO + ByteTrack · 이벤트 DB · 실시간 WebSocket · 관리자 웹  ·  상세: 펫호텔_기능맵_초안.md", 10, GRAY, False)])

    # 08 POSITIONING 2x2
    page += 1
    s = new_slide(prs)
    frame(s, page, "POSITIONING", "영종의 '중대형 × 보이는 안심' 자리는 비어 있습니다")
    ox, oy, aw, ah = 1.6, 1.95, 6.6, 4.5
    box(s, ox + aw / 2, oy, aw / 2, ah / 2, ACCENT_WASH)
    vline(s, ox + aw / 2, oy, ah, RULE)
    hline(s, ox, oy + ah / 2, aw, RULE)
    vline(s, ox, oy, ah, INK, 0.02)
    hline(s, ox, oy + ah, aw, INK, 0.02)
    text(s, ox, oy + ah + 0.12, aw, 0.3, [("수용 범위  →  소형 중심 · · · 중대형·패밀리", 11, GRAY, True)], PP_ALIGN.CENTER)
    text(s, 0.35, oy + ah / 2 - 0.6, 1.15, 1.2, [("안심 투명성", 11, GRAY, True, 2), ("↑ 실시간 캠", 10, GRAY, False, 2), ("· AI 리포트", 10, GRAY, False)])
    players = [
        ("독독", 0.22, 0.48, INK, "공항 공식 채널"),
        ("더우프앤", 0.48, 0.42, INK, "공항 샌딩·장기 할인"),
        ("앤드", 0.75, 0.30, INK, "중대형·생활권"),
        ("호텔독스(서울)", 0.45, 0.85, GRAY, "참고 사례 · 권역 외"),
        ("파우코리아", 0.80, 0.82, ACCENT, "지향하는 위치"),
    ]
    for name, px, py, color, note in players:
        cx = ox + px * aw
        cy = oy + (1 - py) * ah
        dot(s, cx, cy, 0.34 if color == ACCENT else 0.22, color)
        text(s, cx + 0.22, cy - 0.22, 2.2, 0.6, [(name, 13, color, True, 0), (note, 10, GRAY, False)])
    text(
        s,
        8.75,
        2.0,
        3.9,
        4.3,
        [
            ("읽는 법", 13, ACCENT, True, 8),
            ("공식 채널(독독)은 접근성이 좋지만 소형견 중심입니다.", 13, INK, False, 8),
            ("로컬(앤드)은 중대형을 받지만, 공개 정보상 AI 리포트형 안심은 없습니다.", 13, INK, False, 8),
            ("서울에는 AI 캠으로 프리미엄을 만드는 사례가 있습니다. 영종에는 아직 없습니다.", 13, INK, False, 14),
            ("공개 정보 기반의 정성 배치입니다.", 10, GRAY, False),
        ],
    )

    # 09 ASSET
    page += 1
    s = new_slide(prs)
    frame(s, page, "ASSET", f"공사를 마친 1동 4층, {total_rooms}실 규모로 오픈을 준비합니다")
    floors = [
        ("ROOF", "루프탑 가든", "피크닉 존 · 야외 테라스와 정원", False),
        ("3F", "대형 22실 · 패밀리 4실", "VIP 2타입 · 라운지 — 중대형 전용 층", True),
        ("2F", "소형 52실 · 중형 34실", "2타입 객실 · 라운지 · 펫 플레이존", True),
        ("1F", "프런트 · 상담/예약 · 스파 · 미용", "휴식실 · 카페&베이커리 — 체크인과 부가 서비스", True),
    ]
    for i, (fl, name, note, core) in enumerate(floors):
        y = 1.95 + i * 1.12
        box(s, MX, y, 1.1, 0.95, ACCENT if core else SAND)
        text(s, MX, y, 1.1, 0.95, [(fl, 16, WHITE if core else GRAY, True)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        box(s, MX + 1.2, y, 5.6, 0.95, WHITE if core else SAND)
        text(s, MX + 1.45, y + 0.14, 5.2, 0.8, [(name, 15, INK if core else GRAY, True, 2), (note, 11, GRAY, False)])
    text(
        s,
        8.0,
        1.95,
        4.6,
        4.6,
        [
            ("위치", 13, ACCENT, True, 6),
            ("인천 중구 운남동로 44 (영종 운남지구)", 14, INK, True, 4),
            ("공항철도 영종역 5분 · 금산IC 약 3km · 인천공항 약 15분", 12, GRAY, False, 4),
            ("하늘도시 · 미단시티 · 인스파이어 약 2km", 12, GRAY, False, 12),
            ("상태", 13, ACCENT, True, 6),
            ("1차 공사 완료, 오픈 준비 단계", 14, INK, True, 12),
            ("인허가", 13, ACCENT, True, 6),
            ("동물위탁관리업 등록 완료", 14, INK, True, 12),
            ("원 사업계획서(리조이스) 평면 기준 · 컨셉은 진행 중 바뀔 수 있습니다", 10, GRAY, False),
        ],
    )

    # 10 EXPANSION
    page += 1
    s = new_slide(prs)
    frame(s, page, "OPTION VALUE", "2차 개발은 1차 가동 데이터를 보고 결정합니다")
    options = [
        ("B1F", "펫 플레이파크", "약 1,054평 (3,485㎡)", ["데이케어 · 어질리티 존", "펫파크 2곳 · 수영장", "레스토랑 · 레슨룸"]),
        ("나동 1F", "동물병원", "투숙 중 건강 대응", ["진료 · 대기 공간", "위탁견 응급 연계", "건강 기록과 리포트 연결"]),
        ("나동 2–3F", "펫몰 · 파티룸", "보호자 체류 콘텐츠", ["펫 푸드 · 용품 · 패션", "생일 · 모임 파티룸", "쇼룸"]),
    ]
    for i, (tag, name, sub, items) in enumerate(options):
        x = MX + i * 4.05
        hline(s, x, 2.0, 3.75, INK, 0.02)
        text(s, x, 2.2, 3.7, 0.3, [(tag, 12, ACCENT, True)])
        text(s, x, 2.6, 3.7, 0.9, [(name, 22, INK, True, 2), (sub, 12, GRAY, False)])
        text(s, x, 3.75, 3.7, 1.6, [(f"·  {t}", 13, INK, False, 6) for t in items])
    box(s, MX, 5.45, 11.93, 1.05, SAND)
    text(
        s,
        MX + 0.35,
        5.6,
        11.3,
        0.85,
        [
            ("이번 투자 범위는 1차(가동 1동 4층)입니다.", 14, INK, True, 4),
            ("2차는 원 사업계획에도 '추후 일정 확정'으로 되어 있습니다. 1차 가동률과 재이용 데이터로 판단합니다.", 12, GRAY, False),
        ],
    )

    # 11 UNIT ECONOMICS
    page += 1
    s = new_slide(prs)
    frame(s, page, "UNIT ECONOMICS", f"가동률 10%p마다 연매출이 약 {step_10pt:.1f}억 원씩 달라집니다")
    data = CategoryChartData()
    data.categories = [f"가동률 {int(OCC_SCENARIOS[0] * 100)}%"] + [f"{int(o * 100)}%" for o in OCC_SCENARIOS[1:]]
    data.add_series("연매출", [round(r, 1) for r in revenues])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, I(MX), I(1.95), I(6.0), I(4.0), data)
    style_chart(gf.chart, '0.0"억 원"', [SAND, ACCENT, SAND])
    text(s, MX, 6.05, 6.0, 0.6, [("시나리오 연매출(숙박 + 부가 15%) · 실사 전 가정 · 비용 미반영", 11, GRAY, False)])
    room_lines = [
        (f"{name} {n}실 · 1박 {price / 10000:g}만 원(가정)", 13, INK, False, 3) for name, n, price in ROOMS
    ]
    text(
        s,
        7.2,
        1.95,
        5.4,
        4.7,
        [("객실 구성 (원 사업계획서)과 가정 요금", 13, ACCENT, True, 8)]
        + room_lines
        + [
            (f"가중 평균 약 {adr / 10000:.1f}만 원 — 권역 공개가 5.5~11만 원 범위 안", 11, GRAY, False, 12),
            ("가동률을 보수적으로 잡은 이유", 13, ACCENT, True, 6),
            (f"{total_rooms}실은 위탁 전용으로 큰 규모라, 동반형 참고치(키녹 ~79%, 소노펫 ~70%)보다 낮게 가정했습니다.", 12, INK, False, 8),
            ("가동률 60%면 하루 약 67마리로, 법정 기준(20마리당 1명)상 최소 4명 이상 상시 인력이 필요합니다.", 12, INK, False),
        ],
    )

    # 12 BUSINESS MODEL
    page += 1
    s = new_slide(prs)
    frame(s, page, "BUSINESS MODEL", "매출원은 한 번에 하나씩, 순서대로 켭니다")
    engines = [
        ("NOW", "위탁 숙박 + 1층 부가", "4타입 객실 · 체중별 요금\n스파 · 미용 · 카페 · 샌딩", "가동률"),
        ("NEXT", "PET-Aeye 안심 옵션", "실시간 열람 + AI 리포트\n1박당 부가 · VIP 객실 포함", "옵션 선택률 · 재이용"),
        ("LATER", "PET-Aeye SaaS · 2차", "타 펫호텔·유치원 구독\nB1F 데이케어 · 나동", "외부 고객 · 확장 판단"),
    ]
    for i, (tag, name, desc, kpi) in enumerate(engines):
        x = MX + i * 4.05
        top = 2.0 + (2 - i) * 0.35
        hgt = 4.3 - (2 - i) * 0.35
        box(s, x, top, 3.8, hgt, ACCENT if i == 0 else (INK if i == 1 else SAND))
        fg = WHITE if i < 2 else INK
        sub = ACCENT_DIM if i == 0 else (NIGHT_GRAY if i == 1 else GRAY)
        text(s, x + 0.3, top + 0.3, 3.3, 0.3, [(tag, 11, sub, True)])
        text(s, x + 0.3, top + 0.7, 3.3, 0.5, [(name, 20, fg, True)])
        text(s, x + 0.3, top + 1.4, 3.3, 1.4, [(d, 13, fg, False, 3) for d in desc.split("\n")])
        text(s, x + 0.3, top + hgt - 0.6, 3.3, 0.35, [(f"KPI  {kpi}", 12, sub, True)])

    # 13 ROADMAP
    page += 1
    s = new_slide(prs)
    frame(s, page, "ROADMAP", "단계마다 확인 지점을 두고 다음으로 넘어갑니다")
    phases = [
        ("지금", "Phase 0 · 오픈 준비", ["위탁관리업 등록 완료 ✓", "인력 채용 · 운영 매뉴얼", f"{total_rooms}실 카메라 · PET-Aeye 설치"], "확인  전 객실 관제 연결"),
        ("~90일", "Phase 1 · 오픈", ["D-30 사전 마케팅 · 제휴", "그랜드 오픈", "공항 안심 보딩 출시"], "확인  첫 안심 Stay"),
        ("~6개월", "Phase 2 · 정착", ["안심 옵션 요금화", "예약 → Stay 연동", "생활권 멤버십"], "확인  가동률 · 재이용"),
        ("12개월+", "Phase 3 · 확장 검토", ["PET-Aeye 외부 판매", "제휴 시설 도입", "2차 개발 여부 판단"], "확인  외부 고객 1곳"),
    ]
    hline(s, MX, 2.35, 11.93, INK, 0.02)
    for i, (when, name, items, gate) in enumerate(phases):
        x = MX + i * 3.05
        dot(s, x + 0.12, 2.36, 0.24, ACCENT if i <= 1 else INK)
        text(s, x, 1.85, 2.8, 0.35, [(when, 13, ACCENT if i <= 1 else GRAY, True)])
        text(s, x, 2.7, 2.85, 0.45, [(name, 15, INK, True)])
        text(s, x, 3.3, 2.85, 1.8, [(f"·  {t}", 13, INK, False, 6) for t in items])
        box(s, x, 5.3, 2.8, 0.55, SAND)
        text(s, x + 0.18, 5.3, 2.6, 0.55, [(gate, 11, INK, True)], anchor=MSO_ANCHOR.MIDDLE)
    text(
        s,
        MX,
        6.15,
        11.9,
        0.5,
        [("사전 마케팅 제휴 후보(원 사업계획): 여행사 · 카드사 · 숙박 앱 · 국내 호텔 · 반려견 훈련기관 · 관련 학과 · 지자체", 11, GRAY, False)],
    )

    # 14 RISKS
    page += 1
    s = new_slide(prs)
    frame(s, page, "RISKS", "가설이 틀릴 수 있는 지점과 확인 방법")
    cols = [("가설", 0.0, 4.4), ("틀린다면", 4.6, 3.6), ("확인 · 대응", 8.4, 3.5)]
    for name, dx, w in cols:
        text(s, MX + dx, 1.9, w, 0.3, [(name, 12, GRAY, True)])
    hline(s, MX, 2.3, 11.93, INK, 0.02)
    risks = [
        (f"{total_rooms}실을 채울 만큼 수요가 있다", "가동률 장기 저조", "타입별 단계 오픈 · 3F 중대형 우선 판매"),
        ("공항 보호자가 공식 카운터 밖으로도 움직인다", "샌딩 상품 수요 부족", "오픈 90일 공항 유입 비중 측정"),
        ("보호자가 '보이는 안심'에 비용을 더 낸다", "옵션 선택 저조", "오픈 초기 무료 체험 → 유료 전환율 측정"),
        ("AI 이벤트가 운영에 쓸 만큼 정확하다", "오탐 · 과알림", "스태프 확인 후 공유 · 고정확 규칙만 알림"),
        ("24시간 인력 기준(20마리당 1명)을 맞춘다", "인건비 부담", "가동률 연동 인력 운영 · AI 관제로 순찰 보조"),
    ]
    for i, (a, b, c) in enumerate(risks):
        y = 2.5 + i * 0.85
        text(s, MX, y, 4.4, 0.75, [(a, 13, INK, True)])
        text(s, MX + 4.6, y, 3.6, 0.75, [(b, 13, GRAY, False)])
        text(s, MX + 8.4, y, 3.5, 0.75, [(c, 12, INK, False)])
        hline(s, MX, y + 0.72, 11.93)

    # 15 ASK
    page += 1
    s = new_slide(prs)
    frame(s, page, "THE ASK", "투자금은 오픈과 '보이는 안심' 가동에 씁니다")
    box(s, MX, 1.95, 5.2, 4.5, INK)
    text(s, MX + 0.4, 2.25, 4.5, 0.3, [("라운드 · 금액 · 지분", 12, NIGHT_GRAY, True)])
    text(s, MX + 0.4, 2.7, 4.5, 0.9, [("TBD", 54, ACCENT, True)])
    text(
        s,
        MX + 0.4,
        3.85,
        4.5,
        2.4,
        [
            ("자금 용도(안)", 12, NIGHT_GRAY, True, 8),
            ("오픈 준비 · 초기 인력", 14, WHITE, False, 6),
            (f"{total_rooms}실 카메라 · PET-Aeye 설치", 14, WHITE, False, 6),
            ("공항 샌딩 차량 · 사전 마케팅", 14, WHITE, False, 6),
            ("운전자본 · 예비비", 14, WHITE, False, 6),
        ],
    )
    text(s, 6.4, 1.95, 6.2, 0.35, [("투자 후 90일 안에 보여드릴 것", 13, ACCENT, True)])
    proofs = [
        ("그랜드 오픈 · 영업 개시", "위탁관리업 등록 완료 상태"),
        ("PET-Aeye 전 객실 가동", f"{total_rooms}실 × 카메라 연결"),
        ("공항 안심 보딩 첫 판매", "샌딩 포함 패키지"),
        ("PET-Aeye 안심 옵션 판매 개시", "유료 선택률 첫 측정"),
    ]
    for i, (h, b) in enumerate(proofs):
        y = 2.5 + i * 0.95
        text(s, 6.4, y, 0.6, 0.6, [(f"{i + 1}", 26, ACCENT, True)])
        text(s, 7.1, y + 0.05, 5.5, 0.8, [(h, 16, INK, True, 2), (b, 12, GRAY, False)])
        hline(s, 6.4, y + 0.82, 6.2)

    # 16 APPENDIX
    page += 1
    s = new_slide(prs)
    frame(s, page, "APPENDIX", "참고한 자료")
    left = [
        ("시장", 13, ACCENT, True, 6),
        ("Grand View Horizon — South Korea Pet Boarding Services (2024 base)", 12, INK, False, 4),
        ("Mordor Intelligence — Pet Hotels Market (2026–2031)", 12, INK, False, 4),
        ("KB금융지주 경영연구소 — 펫테크 기기 사용 이유", 12, INK, False, 4),
        ("한국관광공사 반려동물 동반여행 현황·인식조사 2024", 12, INK, False, 14),
        ("입지 · 채널", 13, ACCENT, True, 6),
        ("인천국제공항공사 보도자료 2024.3 · 2024.11 (호텔링·펫라운지)", 12, INK, False, 4),
        ("인천공항 영문 안내 — 호텔링 약 13kg 이하", 12, INK, False, 4),
        ("인스파이어 리조트 이용약관·FAQ — 반려동물 입실 불가", 12, INK, False, 4),
    ]
    right = [
        ("시설", 13, ACCENT, True, 6),
        ("원 사업계획서(구칭 리조이스 · PETEL 영종) — 객실 구성, 층별 평면, 2차 계획", 12, INK, False, 14),
        ("경쟁 · 가격", 13, ACCENT, True, 6),
        ("독독 공개 요금(5.5~11만 원/박), 더우프앤·앤드 공개 정보", 12, INK, False, 4),
        ("언론 보도 — 키녹·소노펫 객실점유율(동반형 참고치)", 12, INK, False, 4),
        ("호텔독스 — 실시간 CCTV·AI 리포트 포지셔닝", 12, INK, False, 14),
        ("내부", 13, ACCENT, True, 6),
        ("펫호텔_기능맵_초안.md · 프로젝트_범위_확장_초안.md · PET-Aeye", 12, INK, False, 14),
        ("본 자료는 논의를 위한 초안입니다. 수치는 공개 자료와 가정에 기반하며, 투자 판단에는 별도의 실사와 자문이 필요합니다.", 10, GRAY, False),
    ]
    text(s, MX, 2.0, 5.8, 4.6, left)
    text(s, 6.9, 2.0, 5.8, 4.6, right)

    prs.save(OUT)
    print(f"saved: {OUT}  slides={len(prs.slides)}  rooms={total_rooms}  adr={adr:,.0f}  revenues={[round(r, 1) for r in revenues]}")


if __name__ == "__main__":
    build()
