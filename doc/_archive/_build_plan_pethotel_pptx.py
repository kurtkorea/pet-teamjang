# -*- coding: utf-8 -*-
"""PAWKOREA Pet Hotel Yeongjong × PET-Aeye 기획서 v0.1 — hotel contents, open tasks, PET-Aeye integration, SaaS expansion."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

PAPER = RGBColor(0xF5, 0xF3, 0xEE)
INK = RGBColor(0x14, 0x14, 0x14)
GRAY = RGBColor(0x6B, 0x6B, 0x6B)
RULE = RGBColor(0xD9, 0xD5, 0xCC)
SAND = RGBColor(0xEA, 0xE6, 0xDD)
ACCENT = RGBColor(0xFF, 0x5B, 0x1F)
ACCENT_DIM = RGBColor(0xFF, 0xC9, 0xB3)
ACCENT_WASH = RGBColor(0xFF, 0xEC, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NIGHT = RGBColor(0x14, 0x14, 0x14)
NIGHT_GRAY = RGBColor(0x9A, 0x9A, 0x9A)

W, H = Inches(13.333), Inches(7.5)
MX = 0.7
CW = 11.93
FONT = "Malgun Gothic"
OUT = r"D:\dev_pawkorea\pet-teamjang\doc\PLAN_PAWKOREA_pet_hotel_PET-Aeye_v0.1.pptx"
MOCK = r"D:\dev_pawkorea\pet-teamjang\doc\_mockups"
ADMIN_IMG = MOCK + r"\out_admin_wall.png"
STAFF_IMG = MOCK + r"\out_staff_app.png"
GUARD_IMG = MOCK + r"\out_guardian_app.png"

# 객실 구성: 원 사업계획서(리조이스) 2F·3F 평면 기준, 요금·체중 기준은 가정
ROOMS = [
    ("소형", "2F", 52, "~7kg 소형견", 50_000, "객실당 1대 (다실 1캠 검토)"),
    ("중형", "2F", 34, "7~15kg 중형견", 65_000, "객실당 1대"),
    ("대형", "3F", 22, "15kg 이상 대형견", 90_000, "객실당 1대 · 우선 적용"),
    ("패밀리", "3F", 4, "다견 가족 · VIP", 150_000, "객실당 1~2대"),
]
ADDON = 0.15
OCC_SCENARIOS = (0.30, 0.45, 0.60)


def I(v):
    return Inches(v)


def _font(run, size, color, bold):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, color, bold) or (text, size, color, bold, space_after)."""
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        t, size, color, bold = line[:4]
        after = line[4] if len(line) > 4 else 6
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        r = p.add_run()
        r.text = t
        _font(r, size, color, bold)
    return tb


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def hline(slide, x, y, w, color=RULE, weight=0.015):
    return box(slide, x, y, w, weight, color)


def vline(slide, x, y, h, color=RULE, weight=0.015):
    return box(slide, x, y, weight, h, color)


def dot(slide, cx, cy, d, color):
    return box(slide, cx - d / 2, cy - d / 2, d, d, color, MSO_SHAPE.OVAL)


def new_slide(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = NIGHT if dark else PAPER
    return s


def frame(slide, page, section, title, dark=False):
    fg = WHITE if dark else INK
    sub = NIGHT_GRAY if dark else GRAY
    text(slide, MX, 0.45, 8, 0.3, [(section, 11, ACCENT, True)])
    text(slide, MX, 0.8, 11.9, 1.0, [(title, 25, fg, True)])
    text(slide, MX, 7.0, 8, 0.25, [("PAWKOREA Pet Hotel Yeongjong × PET-Aeye  ·  기획서", 9, sub, False)])
    text(slide, 11.63, 7.0, 1.0, 0.25, [(f"{page:02d}", 9, sub, True)], PP_ALIGN.RIGHT)


def table(slide, x, y, cols, rows, row_h, body=12, strong_first=True, head_gap=0.45):
    """Editorial table: header labels, a heavy rule, then rows separated by thin rules. Cells may contain '\\n'."""
    cx = x
    for name, w in cols:
        text(slide, cx, y, w - 0.15, 0.3, [(name, 11, GRAY, True)])
        cx += w
    total_w = sum(w for _, w in cols)
    hline(slide, x, y + head_gap - 0.08, total_w, INK, 0.02)
    ry = y + head_gap + 0.05
    for row in rows:
        cx = x
        for ci, ((_, w), cell) in enumerate(zip(cols, row)):
            bold = strong_first and ci == 0
            color = INK if (ci == 0 or not strong_first) else INK
            text(slide, cx, ry, w - 0.15, row_h - 0.05, [(t, body, color, bold, 1) for t in cell.split("\n")])
            cx += w
        hline(slide, x, ry + row_h - 0.1, total_w)
        ry += row_h
    return ry


def room_economics():
    total = sum(r[2] for r in ROOMS)
    adr = sum(r[2] * r[4] for r in ROOMS) / total
    full_year = total * 365 * adr * (1 + ADDON)
    revenues = [full_year * occ / 1e8 for occ in OCC_SCENARIOS]
    return total, adr, revenues


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    total_rooms, adr, revenues = room_economics()
    page = 0

    # 01 COVER
    page += 1
    s = new_slide(prs)
    box(s, 0, 0, 13.333, 0.12, ACCENT)
    text(s, MX, 0.7, 9, 0.3, [("PLANNING DOC  ·  2026.10  ·  v0.1  ·  내부 논의용", 11, GRAY, True)])
    text(
        s,
        MX,
        1.55,
        11.8,
        2.6,
        [
            ("파우코리아 펫호텔 영종", 46, INK, True, 0),
            ("× PET-Aeye 기획서", 46, ACCENT, True, 0),
        ],
    )
    text(
        s,
        MX,
        4.0,
        11,
        0.9,
        [
            ("호텔에 담을 것 · 오픈까지 할 일 · PET-Aeye 연계 · SaaS 확장", 18, INK, True, 4),
            ("맡기는 순간부터, 보이는 안심을. 영종에서 운영하고, 다른 펫호텔로 넓힙니다.", 14, GRAY, False),
        ],
    )
    hline(s, MX, 5.55, CW, INK, 0.02)
    facts = [
        (f"{total_rooms}실", "4개 타입 위탁 객실"),
        ("1동 4층", "공사 완료 · 오픈 준비"),
        ("PET-Aeye", "실시간 AI 관제 솔루션"),
        ("SaaS", "타 펫호텔 확장"),
    ]
    for i, (v, lab) in enumerate(facts):
        text(s, MX + i * 3.0, 5.8, 2.8, 0.5, [(v, 24, INK, True, 2), (lab, 12, GRAY, False)])

    # 02 OVERVIEW
    page += 1
    s = new_slide(prs)
    frame(s, page, "OVERVIEW", "이 기획서가 다루는 것")
    text(
        s,
        MX,
        1.95,
        5.4,
        4.8,
        [
            ("목적", 13, ACCENT, True, 6),
            ("파우코리아 펫호텔 영종의 오픈·운영 계획과, 자체 AI 관제 PET-Aeye의 호텔 적용 · SaaS 확장 계획을 한 문서로 정리합니다.", 14, INK, False, 16),
            ("범위", 13, ACCENT, True, 6),
            (f"위탁형 펫호텔 1차(1동 4층 · {total_rooms}실) · PET-Aeye 호텔 적용 · PET-Aeye SaaS", 14, INK, False, 4),
            ("제외: 동반숙박 · 장례 · 2차 개발(B1F·나동) 상세", 12, GRAY, False, 16),
            ("전제", 13, ACCENT, True, 6),
            ("동물위탁관리업 등록 완료 · 1차 공사 완료 · PET-Aeye 관제 솔루션 운영", 14, INK, False),
        ],
    )
    vline(s, 6.55, 1.95, 4.7)
    parts = [
        ("01", "호텔에 담을 것", "공간 · 객실 · 서비스 · 운영 기준"),
        ("02", "해야 할 일", "오픈 준비 워크스트림 · 일정"),
        ("03", "PET-Aeye 연계", "여정 · 카메라 · 이벤트 · 데이터"),
        ("04", "SaaS 확장", "고객 · 패키지 · 제품 요건 · 로드맵 · 채널"),
        ("05", "관리", "지표 · 리스크 · 결정 사항"),
    ]
    for i, (n, h, d) in enumerate(parts):
        y = 1.95 + i * 0.92
        text(s, 6.9, y, 0.8, 0.6, [(n, 22, ACCENT, True)])
        text(s, 7.75, y + 0.02, 4.9, 0.8, [(h, 16, INK, True, 2), (d, 12, GRAY, False)])

    # 03 STRUCTURE
    page += 1
    s = new_slide(prs)
    frame(s, page, "OVERVIEW", "호텔 · PET-Aeye · SaaS는 하나의 흐름입니다")
    layers = [
        ("HOTEL", "호텔 운영", [f"{total_rooms}실 위탁 숙박", "케어 · 스파 · 미용 · 샌딩", "현장 운영 데이터"], "매출 · 현장 검증", ACCENT),
        ("PET-Aeye", "AI 관제 · 안심", ["보고: 룸별 라이브 관제", "알리고: 이상 행동 알림", "요약: 일일 CareReport"], "차별화 · 운영 효율", INK),
        ("SaaS", "타 시설 공급", ["위탁형 펫호텔", "유치원 · 데이케어", "병원 입원실 · 훈련소"], "두 번째 매출원", SAND),
    ]
    for i, (tag, name, items, role, color) in enumerate(layers):
        x = MX + i * 4.1
        box(s, x, 1.95, 3.6, 3.3, color)
        fg = INK if color == SAND else WHITE
        sub = GRAY if color == SAND else (ACCENT_DIM if color == ACCENT else NIGHT_GRAY)
        text(s, x + 0.3, 2.2, 3.1, 0.3, [(tag, 11, sub, True)])
        text(s, x + 0.3, 2.55, 3.1, 0.5, [(name, 20, fg, True)])
        text(s, x + 0.3, 3.25, 3.1, 1.4, [(f"·  {t}", 13, fg, False, 4) for t in items])
        text(s, x + 0.3, 4.75, 3.1, 0.35, [(f"역할  {role}", 12, sub, True)])
        if i < 2:
            text(s, x + 3.6, 3.3, 0.5, 0.5, [("→", 22, GRAY, True)], PP_ALIGN.CENTER)
    box(s, MX, 5.55, CW, 1.0, SAND)
    text(
        s,
        MX + 0.35,
        5.7,
        11.3,
        0.8,
        [
            ("호텔은 PET-Aeye의 첫 고객이고, PET-Aeye는 호텔의 차별점입니다.", 15, INK, True, 4),
            ("영종에서 다듬은 운영 방식을 제품에 그대로 담아 다른 시설에 공급합니다.", 13, GRAY, False),
        ],
    )

    # 04 SPACE
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 1 · 호텔에 담을 것", "층별 공간과 추가로 갖출 운영 공간")
    floors = [
        ("ROOF", "루프탑 가든", "피크닉 존 · 야외 테라스 · 정원", False),
        ("3F", "대형 22실 · 패밀리 4실", "VIP 2타입 · 라운지 — 중대형 전용 층", True),
        ("2F", "소형 52실 · 중형 34실", "2타입 객실 · 라운지 · 펫 플레이존", True),
        ("1F", "프런트 · 상담/예약 · 스파 · 미용", "휴식실 · 카페&베이커리", True),
    ]
    for i, (fl, name, note, core) in enumerate(floors):
        y = 1.95 + i * 1.12
        box(s, MX, y, 1.1, 0.95, ACCENT if core else SAND)
        text(s, MX, y, 1.1, 0.95, [(fl, 16, WHITE if core else GRAY, True)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        box(s, MX + 1.2, y, 5.3, 0.95, WHITE if core else SAND)
        text(s, MX + 1.45, y + 0.14, 4.9, 0.8, [(name, 15, INK if core else GRAY, True, 2), (note, 11, GRAY, False)])
    text(s, MX, 6.45, 6.5, 0.3, [("원 사업계획서(리조이스) 평면 기준", 10, GRAY, False)])
    text(s, 7.7, 1.95, 4.9, 0.35, [("추가로 갖출 운영 공간", 13, ACCENT, True)])
    extra = [
        ("관제실", "PET-Aeye 룸 그리드 모니터 · 야간 관제"),
        ("격리실", "입실 후 이상 증상 개체 분리"),
        ("급식 준비실", "개체별 식단 · 특식 · 투약 준비"),
        ("세탁 · 소독실", "객실 회전 청소 · 침구 세탁"),
        ("비품 · 의약품 보관", "응급 키트 · 투약 기록"),
        ("이중문 · 출입 통제", "층별 탈출 방지 · 출입 기록"),
    ]
    for i, (h, d) in enumerate(extra):
        y = 2.45 + i * 0.68
        text(s, 7.7, y, 1.9, 0.5, [(h, 13, INK, True)])
        text(s, 9.65, y, 3.0, 0.5, [(d, 12, GRAY, False)])
        hline(s, 7.7, y + 0.5, 4.93)
    text(s, 7.7, 6.45, 4.9, 0.3, [("위치는 현장 실측 후 확정합니다.", 10, GRAY, False)])

    # 05 ROOMS
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 1 · 호텔에 담을 것", "객실 상품 — 4개 타입, 112실")
    cols = [("타입", 1.5), ("위치", 1.0), ("실수", 1.0), ("대상 (가정)", 2.6), ("1박 요금 (가정)", 2.2), ("카메라", 3.63)]
    rows = [(n, fl, f"{c}실", tgt, f"{p / 10000:g}만 원", cam) for n, fl, c, tgt, p, cam in ROOMS]
    rows.append(("합계", "", f"{total_rooms}실", "", f"평균 약 {adr / 10000:.1f}만 원", "전 객실 PET-Aeye 연결"))
    table(s, MX, 1.95, cols, rows, 0.62, body=13)
    box(s, MX, 5.4, CW, 1.15, SAND)
    text(
        s,
        MX + 0.35,
        5.55,
        11.3,
        0.95,
        [
            ("3F 26실(대형 · 패밀리)을 중대형 수요의 핵심 상품으로 둡니다.", 14, INK, True, 4),
            (
                f"인천공항 공식 호텔링은 약 13kg 이하 중심입니다. 가동률 30 · 45 · 60% 기준 연매출은 약 "
                f"{revenues[0]:.1f} · {revenues[1]:.1f} · {revenues[2]:.1f}억 원입니다(부가 15% 포함, 가정).",
                12,
                GRAY,
                False,
            ),
        ],
    )

    # 06 SERVICES
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 1 · 호텔에 담을 것", "서비스 메뉴 — 기본 · 부가 · 안심")
    menus = [
        ("BASIC", "기본 위탁", "모든 Stay", ["개체별 급식 · 급수", "하루 2회 이상 산책 · 놀이", "배변 처리 · 객실 청소", "24시간 상주 케어", "일일 케어 로그"]),
        ("ADD-ON", "부가 서비스", "선택", ["펫 스파 · 목욕", "미용", "공항 샌딩 · 픽업", "특식 (SPECIAL PET FOOD)", "투약 관리 · 데이케어"]),
        ("PET-Aeye", "안심 서비스", "PET-Aeye 제공", ["배정 룸 라이브 열람", "일일 CareReport", "이상 시 보호자 알림 (선택)", "퇴실 하이라이트 클립", "Stay 이력 · 성향 기록"]),
    ]
    for i, (tag, name, sub, items) in enumerate(menus):
        x = MX + i * 4.05
        hline(s, x, 2.0, 3.75, ACCENT if i == 2 else INK, 0.03)
        text(s, x, 2.2, 3.7, 0.3, [(tag, 11, ACCENT, True)])
        text(s, x, 2.55, 3.7, 0.8, [(name, 21, INK, True, 2), (sub, 12, GRAY, False)])
        text(s, x, 3.55, 3.7, 2.0, [(f"·  {t}", 13, INK, False, 6) for t in items])
    box(s, MX, 5.6, CW, 0.95, ACCENT_WASH)
    text(
        s,
        MX + 0.35,
        5.72,
        11.3,
        0.8,
        [
            ("대표 패키지  ‘공항 안심 보딩’ = 숙박 + 공항 샌딩 · 픽업 + 안심 서비스", 14, INK, True, 4),
            ("인스파이어 투숙객 숏스테이 · 영종 생활권 데이케어는 같은 구성에서 기간만 다르게 운영합니다.", 12, GRAY, False),
        ],
    )

    # 07 OPERATING STANDARDS
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 1 · 호텔에 담을 것", "운영 기준 — 입실 · 케어 · 인력 · 정책")
    quads = [
        ("입실 기준", ["예방접종 증명 확인", "건강 문진 · 기저질환 · 복용약", "공격성 · 분리불안 체크", "생리 · 발정 여부"]),
        ("케어 기준", ["개체별 식단 · 급여 시간", "산책 · 놀이 일정", "순찰과 PET-Aeye 관제 병행", "케어 로그 기록 · 리포트 반영"]),
        ("인력 기준", ["24시간 상주", "위탁 20마리당 1명 이상", "가동률 연동 근무표", "야간 관제 담당 지정"]),
        ("정책 · 동의", ["이용 약관 · 취소 / 환불", "응급 진료 동의", "영상 촬영 · 열람 · 보존 고지", "배상책임보험 가입"]),
    ]
    for i, (h, items) in enumerate(quads):
        cx = MX + (i % 2) * 6.1
        cy = 1.95 + (i // 2) * 2.35
        hline(s, cx, cy, 5.8, INK, 0.02)
        text(s, cx, cy + 0.18, 5.7, 0.4, [(h, 16, ACCENT if i == 2 else INK, True)])
        text(s, cx, cy + 0.7, 5.7, 1.6, [(f"·  {t}", 13, INK, False, 5) for t in items])
    text(s, MX, 6.6, 11.9, 0.3, [("가동률 60%(하루 약 67마리)면 상시 최소 4명이 필요합니다. 근무표는 객실 판매 현황과 함께 짭니다.", 11, GRAY, False)])

    # 08 OPEN WORKSTREAMS
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 2 · 해야 할 일", "오픈 준비 — 8개 워크스트림")
    cols = [("영역", 2.0), ("할 일", 7.33), ("시작", 1.2), ("담당", 1.4)]
    rows = [
        ("시설 · 인프라", f"{total_rooms}실 · 복도 카메라 설치, 네트워크 · 녹화 스토리지, 관제실 구성", "D-60", "시설"),
        ("인력 · 교육", "케어 · 프런트 · 미용 · 드라이버 채용, PET-Aeye · 응급 대응 교육", "D-60", "운영"),
        ("운영 매뉴얼", "체크인 / 아웃 · 케어 · 응급 · 이벤트 대응 SOP 작성", "D-45", "운영"),
        ("상품 · 가격", "객실 요금표, 안심 서비스 요금 방식, 공항 패키지 확정", "D-45", "기획"),
        ("정책 · 법무", "이용 약관, 동의서, 영상 개인정보 고지, 배상책임보험", "D-45", "기획"),
        ("예약 · 결제", "예약 채널 개설, 예약 → PET-Aeye Stay 연동", "D-30", "기획 · 개발"),
        ("PET-Aeye 세팅", "룸 · 카메라 매핑, 알림 규칙 조정, 리포트 템플릿, 리허설", "D-21", "개발"),
        ("마케팅 · 제휴", "사전 마케팅, 여행사 · 카드사 · 숙박 앱 · 호텔 제휴", "D-30", "마케팅"),
    ]
    table(s, MX, 1.9, cols, rows, 0.56, body=12)

    # 09 TIMELINE
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 2 · 해야 할 일", "오픈 전 60일, 오픈 후 90일 일정")
    marks = [
        ("D-60", "착수", "카메라 · 네트워크 설치\n인력 채용 시작"),
        ("D-45", "기준 확정", "SOP · 요금 · 약관\n동의서 확정"),
        ("D-30", "예약 오픈", "예약 채널 · Stay 연동\n사전 마케팅"),
        ("D-14", "리허설", "전 객실 관제 점검\n스태프 교육 완료"),
        ("D-0", "그랜드 오픈", "공항 안심 보딩\n판매 시작"),
        ("D+30", "첫 리뷰", "가동률 · 알림 품질\nSOP 수정"),
        ("D+90", "정착", "안심 서비스 요금 확정\nSaaS 파일럿 후보 선정"),
    ]
    y = 3.0
    hline(s, MX + 0.4, y, 11.1, INK, 0.02)
    gap = 11.1 / (len(marks) - 1)
    for i, (when, what, desc) in enumerate(marks):
        cx = MX + 0.4 + i * gap
        hot = i == 4
        dot(s, cx, y + 0.01, 0.42 if hot else 0.24, ACCENT if hot else INK)
        text(s, cx - 0.9, 2.15, 1.8, 0.4, [(when, 13, ACCENT if hot else GRAY, True)], PP_ALIGN.CENTER)
        text(s, cx - 0.9, 3.45, 1.8, 0.45, [(what, 15, INK, True)], PP_ALIGN.CENTER)
        text(s, cx - 0.9, 3.95, 1.8, 1.0, [(d, 11, GRAY, False, 2) for d in desc.split("\n")], PP_ALIGN.CENTER)
    box(s, MX, 5.35, CW, 1.2, SAND)
    text(
        s,
        MX + 0.35,
        5.52,
        11.3,
        0.95,
        [
            ("오픈 조건: 전 객실 카메라 연결 · SOP 교육 완료 · 예약 → Stay 연동 확인", 14, INK, True, 4),
            ("D-14 리허설에서 세 가지를 모두 확인한 뒤 오픈합니다.", 12, GRAY, False),
        ],
    )

    # 10 PET-Aeye OVERVIEW
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 3 · PET-Aeye 연계", "PET-Aeye — 보고, 알리고, 요약합니다")
    loop = [
        ("보고", "룸별 카메라 라이브\n검출 오버레이"),
        ("알리고", "장시간 무움직임 · 이상 행동\n스태프 우선 알림"),
        ("요약", "Stay 단위 일일 리포트\n보호자 기간 한정 링크"),
    ]
    for i, (h, b) in enumerate(loop):
        x = MX + i * 2.05
        box(s, x, 2.0, 1.85, 1.85, ACCENT if i == 0 else INK, MSO_SHAPE.OVAL)
        text(s, x, 2.0, 1.85, 1.85, [(h, 20, WHITE, True)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(s, x - 0.1, 4.05, 2.05, 1.0, [(t, 11, GRAY, False, 2) for t in b.split("\n")], PP_ALIGN.CENTER)
    text(s, MX, 5.05, 6.1, 1.0, [("원 사업계획의 '24시간 CHECKING SYSTEM'을 사람의 순찰과 PET-Aeye가 함께 맡습니다.", 12, INK, False)])
    vline(s, 7.0, 1.95, 4.5)
    text(s, 7.35, 1.95, 5.3, 0.35, [("스태프용 관제", 13, ACCENT, True)])
    staff = [
        f"{total_rooms}실 룸 그리드 · 이벤트 타임라인",
        "AI 검출 · 추적 · 행동 라벨",
        "이상 행동 알림 · PTZ 자동 추적 · 녹화",
    ]
    text(s, 7.35, 2.35, 5.3, 1.6, [(f"✓  {t}", 13, INK, False, 5) for t in staff])
    text(s, 7.35, 3.95, 5.3, 0.35, [("보호자용 안심", 13, ACCENT, True)])
    guardian = ["기간 한정 링크로 배정 룸 라이브", "Stay 단위 일일 CareReport", "퇴실 후 하이라이트 클립"]
    text(s, 7.35, 4.35, 5.3, 1.6, [(f"✓  {t}", 13, INK, False, 5) for t in guardian])
    text(s, MX, 6.45, 11.9, 0.35, [("구성: RTSP 수신 · 녹화 · YOLO + ByteTrack · 이벤트 DB · 실시간 WebSocket · 관리자 웹 · 보호자 웹", 10, GRAY, False)])

    # 10b ADMIN MONITORING WALL (mockup)
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 3 · PET-Aeye 연계", "관리자 관제 화면 — 112실을 한 화면에서")
    pic = s.shapes.add_picture(ADMIN_IMG, I(MX), I(2.0), height=I(4.6))
    pic.line.color.rgb = RULE
    pic.line.width = Pt(0.75)
    ann = [
        ("룸 그리드 관제", "전 객실 라이브를 한 화면에서 봅니다"),
        ("AI 검출 오버레이", "개체 · 행동 · 신뢰도를 실시간 표시"),
        ("이상 감지 강조", "확인 필요한 객실만 눈에 띄게"),
        ("실시간 이벤트", "급식 · 산책 · 알림 로그 자동 기록"),
    ]
    for i, (h, d) in enumerate(ann):
        y = 2.05 + i * 1.05
        text(s, 9.15, y, 3.45, 0.9, [(h, 15, ACCENT, True, 3), (d, 12, GRAY, False)])
        if i < len(ann) - 1:
            hline(s, 9.15, y + 0.9, 3.45)
    text(s, MX, 6.75, 11.9, 0.3, [("* 화면은 구성 예시입니다.", 10, GRAY, False)])

    # 10c STAFF & GUARDIAN APPS (mockup)
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 3 · PET-Aeye 연계", "스태프 앱 · 보호자 앱")
    sp = s.shapes.add_picture(STAFF_IMG, I(0.95), I(1.75), height=I(5.1))
    gp = s.shapes.add_picture(GUARD_IMG, I(6.95), I(1.75), height=I(5.1))
    text(
        s,
        3.55,
        2.0,
        2.95,
        4.6,
        [
            ("스태프 앱", 16, ACCENT, True, 10),
            ("확인 필요 알림을 가장 위에", 13, INK, True, 2),
            ("한 객실만 빠르게 판단합니다", 11, GRAY, False, 12),
            ("객실 라이브 + AI 오버레이", 13, INK, True, 2),
            ("알림 화면에서 바로 확인", 11, GRAY, False, 12),
            ("현장 확인 / 정상 처리", 13, INK, True, 2),
            ("한 번의 조작으로 처리", 11, GRAY, False, 12),
            ("담당 층 할 일 · 케어 로그", 13, INK, True, 2),
            ("급식 · 산책 · 투약 체크", 11, GRAY, False),
        ],
    )
    text(
        s,
        9.55,
        2.0,
        3.05,
        4.6,
        [
            ("보호자 앱", 16, ACCENT, True, 10),
            ("배정 객실만 실시간 열람", 13, INK, True, 2),
            ("전체 시설이 아닌 내 반려견 방", 11, GRAY, False, 12),
            ("오늘의 CareReport", 13, INK, True, 2),
            ("급식 · 산책 · 휴식 요약", 11, GRAY, False, 12),
            ("하이라이트 클립", 13, INK, True, 2),
            ("하루 활동을 짧은 영상으로", 11, GRAY, False, 12),
            ("기간 한정 안심 링크", 13, INK, True, 2),
            ("체크아웃까지 열리고 자동 만료", 11, GRAY, False),
        ],
    )

    # 11 JOURNEY × PET-Aeye
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 3 · PET-Aeye 연계", "고객 여정마다 PET-Aeye가 붙는 지점")
    steps = [
        ("예약", "룸 타입 · 기간\n옵션 접수", "Stay 사전 생성"),
        ("체크인", "건강 문진\n동의서 서명", "룸 · 카메라 배정\n보호자 링크 발송"),
        ("투숙 중", "급식 · 산책\n청소 · 순찰", "룸 그리드 관제\n이상 행동 알림"),
        ("매일 저녁", "케어 메모\n작성", "CareReport\n자동 생성 · 발송"),
        ("체크아웃", "인계 · 결제", "링크 만료\n하이라이트 클립"),
        ("이후", "재방문 안내", "Stay 이력\n성향 기록"),
    ]
    lx, colw = MX + 1.45, (CW - 1.45) / len(steps)
    text(s, MX, 2.75, 1.3, 0.4, [("호텔 업무", 12, GRAY, True)])
    text(s, MX, 4.35, 1.3, 0.4, [("PET-Aeye", 12, ACCENT, True)])
    hline(s, MX, 2.45, CW, INK, 0.02)
    hline(s, MX, 4.1, CW)
    box(s, lx, 4.15, CW - 1.45, 1.5, ACCENT_WASH)
    for i, (name, hotel, pt) in enumerate(steps):
        x = lx + i * colw
        text(s, x + 0.1, 2.0, colw - 0.2, 0.4, [(name, 15, INK, True)])
        text(s, x + 0.1, 2.75, colw - 0.2, 1.2, [(t, 12, INK, False, 2) for t in hotel.split("\n")])
        text(s, x + 0.1, 4.35, colw - 0.2, 1.2, [(t, 12, INK, True, 2) for t in pt.split("\n")])
        if i:
            vline(s, x, 2.6, 3.05)
    text(s, MX, 6.0, 11.9, 0.6, [("모든 기록은 Stay(위탁 1건)에 묶여, 호텔 업무 단위와 PET-Aeye 데이터 단위가 같아집니다.", 12, GRAY, False)])

    # 12 CAMERA LAYOUT
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 3 · PET-Aeye 연계", "카메라 배치 기준")
    cols = [("구역", 2.3), ("층", 0.9), ("수량 기준", 2.6), ("목적", 2.9), ("보호자 노출", 3.23)]
    rows = [
        ("소형 · 중형 객실", "2F", "86실 · 객실당 1대", "위탁 본체 관제", "배정 룸 라이브"),
        ("대형 · 패밀리 객실", "3F", "26실 · 객실당 1대 이상", "중대형 관제 · 우선 설치", "배정 룸 라이브"),
        ("라운지 · 플레이존", "2–3F", "구역당 1대 (PTZ)", "놀이 시간 관찰", "하이라이트 클립"),
        ("복도 · 이중문", "2–3F", "출입구별 1대", "탈출 방지 · 보안", "비노출"),
        ("프런트 · 스파 · 미용", "1F", "구역당 1대", "체크인 · 서비스 기록", "선택 노출"),
    ]
    table(s, MX, 1.9, cols, rows, 0.62, body=13)
    text(
        s,
        MX,
        5.7,
        11.9,
        1.2,
        [
            ("함께 갖출 인프라", 13, ACCENT, True, 6),
            ("PoE 스위치 · 녹화 스토리지(보존 기간에 맞춰 용량 산정) · 관제실 모니터 · 이중화된 인터넷 회선 · 정전 대비 UPS", 13, INK, False, 4),
            ("최종 수량은 현장 실측과 화각 테스트 후 확정합니다.", 11, GRAY, False),
        ],
    )

    # 13 EVENT → ACTION
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 3 · PET-Aeye 연계", "이벤트가 생기면 누가 무엇을 하는가")
    cols = [("PET-Aeye 신호", 3.4), ("스태프 액션", 4.4), ("보호자 노출", 4.13)]
    rows = [
        ("장시간 무움직임", "현장 확인 · 필요 시 병원 연계", "리포트 기재 · 선택 시 알림"),
        ("과도한 반복 이동", "산책 · 놀이 추가 · 성향 기록", "일일 리포트"),
        ("비인가 인원 검출", "보안 점검 · 출입 기록 확인", "비노출"),
        ("카메라 연결 끊김", "즉시 점검 · 대체 순찰", "비노출"),
        ("Stay 종료", "링크 만료 · 영상 보존 정책 적용", "하이라이트 제공"),
    ]
    table(s, MX, 1.9, cols, rows, 0.62, body=13)
    box(s, MX, 5.45, CW, 1.1, SAND)
    text(
        s,
        MX + 0.35,
        5.6,
        11.3,
        0.9,
        [
            ("알림은 정확도가 확인된 규칙부터 켭니다.", 14, INK, True, 4),
            ("보호자에게는 스태프가 확인한 내용만 공유합니다. 규칙과 기준값은 오픈 후 30일 데이터로 조정합니다.", 12, GRAY, False),
        ],
    )

    # 14 DATA
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 3 · PET-Aeye 연계", "데이터는 Stay 하나로 연결합니다")
    chain_a = [("Guardian · Pet", "보호자 · 반려견"), ("Stay", "위탁 1건"), ("Room", "객실"), ("Camera", "카메라")]
    chain_b = [("Camera", "영상 · 검출"), ("PetEvent", "행동 · 알림"), ("CareReport", "일일 요약"), ("Guardian Link", "보호자 열람")]
    bw, bg = 1.38, 0.32
    for row, chain in enumerate((chain_a, chain_b)):
        y = 2.1 + row * 1.6
        for i, (name, sub) in enumerate(chain):
            x = MX + i * (bw + bg)
            hot = name in ("Stay", "CareReport")
            box(s, x, y, bw, 1.0, ACCENT if hot else WHITE)
            text(
                s,
                x,
                y,
                bw,
                1.0,
                [(name, 12, WHITE if hot else INK, True, 2), (sub, 10, ACCENT_DIM if hot else GRAY, False)],
                PP_ALIGN.CENTER,
                MSO_ANCHOR.MIDDLE,
            )
            if i < len(chain) - 1:
                text(s, x + bw, y + 0.3, bg, 0.4, [("→", 14, GRAY, True)], PP_ALIGN.CENTER)
    text(
        s,
        MX,
        5.0,
        6.3,
        1.5,
        [
            ("체크인 때 스태프가 반려견과 카메라 트랙을 연결합니다.", 13, INK, False, 4),
            ("이후 이벤트 · 리포트 · 열람 기록이 모두 같은 Stay에 쌓입니다.", 13, INK, False),
        ],
    )
    vline(s, 7.35, 1.95, 4.6)
    text(s, 7.7, 1.95, 4.9, 0.35, [("호텔 시스템과 엮는 연동 작업", 13, ACCENT, True)])
    links = [
        ("예약 → Stay", "예약 채널 확정 건으로 Stay 자동 생성"),
        ("케어 체크리스트", "급식 · 산책 · 배변 기록을 이벤트와 함께 저장"),
        ("구역(Zone) 규칙", "급식대 · 문 앞 등 구역별 알림"),
        ("열람 감사 로그", "보호자 · 스태프 영상 열람 기록"),
        ("결제 → 옵션", "안심 서비스 결제 시 보호자 링크 활성화"),
    ]
    for i, (h, d) in enumerate(links):
        y = 2.45 + i * 0.8
        text(s, 7.7, y, 4.9, 0.7, [(h, 13, INK, True, 2), (d, 12, GRAY, False)])
        hline(s, 7.7, y + 0.68, 4.93)

    # 15 SAAS STATEMENT (dark)
    page += 1
    s = new_slide(prs, dark=True)
    text(s, MX, 0.45, 8, 0.3, [("PART 4 · SaaS 확장", 11, ACCENT, True)])
    text(
        s,
        MX,
        1.7,
        11.8,
        2.2,
        [
            ("영종은 PET-Aeye의 첫 고객이자 쇼룸입니다.", 38, WHITE, True, 8),
            ("호텔에서 다듬은 운영 방식을 제품으로 만들어 다른 시설에 공급합니다.", 22, ACCENT, True),
        ],
    )
    flywheel = ["영종 운영", "운영 데이터", "제품 개선", "운영 사례", "외부 판매"]
    fw = 2.0
    for i, step in enumerate(flywheel):
        x = MX + i * (fw + 0.48)
        box(s, x, 4.5, fw, 0.8, ACCENT if i in (0, 4) else RGBColor(0x2A, 0x2A, 0x2A))
        text(s, x, 4.5, fw, 0.8, [(step, 15, WHITE, True)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if i < len(flywheel) - 1:
            text(s, x + fw, 4.68, 0.48, 0.4, [("→", 16, NIGHT_GRAY, True)], PP_ALIGN.CENTER)
    text(
        s,
        MX,
        5.75,
        11.8,
        0.8,
        [("외부 고객이 쓰는 기능은 영종 현장에서 먼저 씁니다. 영종에서 쓰지 않는 기능은 팔지 않습니다.", 13, NIGHT_GRAY, False)],
    )
    text(s, 11.63, 7.0, 1.0, 0.25, [(f"{page:02d}", 9, NIGHT_GRAY, True)], PP_ALIGN.RIGHT)

    # 16 TARGET CUSTOMERS
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 4 · SaaS 확장", "타깃 고객 — 반려견을 맡아 두는 모든 시설")
    segs = [
        ("01 · 핵심", "위탁형 펫호텔", "보호자 문의 응대 부담\n야간 관제 인력 부족", "Plus"),
        ("02", "유치원 · 데이케어", "등원 중 보호자 확인 요구\n놀이 · 사회화 기록", "Plus"),
        ("03", "동물병원 입원실", "입원견 상태 관찰\n보호자 면회 대체", "Basic → Plus"),
        ("04", "위탁 훈련소", "장기 위탁 중\n훈련 진행 공유", "Plus"),
    ]
    for i, (tag, name, need, pkg) in enumerate(segs):
        x = MX + i * 3.03
        hline(s, x, 2.0, 2.8, ACCENT if i == 0 else INK, 0.03)
        text(s, x, 2.2, 2.8, 0.3, [(tag, 11, ACCENT, True)])
        text(s, x, 2.55, 2.8, 0.5, [(name, 18, INK, True)])
        text(s, x, 3.3, 2.8, 0.3, [("필요", 11, GRAY, True)])
        text(s, x, 3.65, 2.8, 1.0, [(t, 13, INK, False, 3) for t in need.split("\n")])
        text(s, x, 4.75, 2.8, 0.3, [("추천 패키지", 11, GRAY, True)])
        text(s, x, 5.1, 2.8, 0.4, [(pkg, 15, ACCENT if i == 0 else INK, True)])
    text(s, MX, 6.2, 11.9, 0.5, [("첫 판매 대상은 영종과 운영 방식이 같은 위탁형 펫호텔입니다. 나머지는 파일럿 결과를 보고 순서를 정합니다.", 12, GRAY, False)])

    # 17 PACKAGES
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 4 · SaaS 확장", "패키지 구성 — Basic · Plus · Pro")
    pkgs = [
        ("BASIC", "스태프 관제", ["카메라 등록 · 룸 매핑", "룸 그리드 라이브", "AI 검출 · 이벤트 타임라인", "녹화 · 보관", "스태프 알림"], SAND),
        ("PLUS", "보호자 안심", ["Basic 전체", "Stay 관리", "보호자 기간 한정 링크", "일일 CareReport · 하이라이트", "시설 로고 · 이름 표시"], ACCENT),
        ("PRO", "다지점 · 연동", ["Plus 전체", "다지점 통합 관제", "예약 시스템 API · 웹훅", "권한 · 열람 감사 로그", "전담 온보딩"], INK),
    ]
    for i, (tag, name, items, color) in enumerate(pkgs):
        x = MX + i * 4.05
        box(s, x, 1.95, 3.8, 3.75, color)
        fg = INK if color == SAND else WHITE
        sub = GRAY if color == SAND else (ACCENT_DIM if color == ACCENT else NIGHT_GRAY)
        text(s, x + 0.3, 2.2, 3.3, 0.3, [(tag, 11, sub, True)])
        text(s, x + 0.3, 2.55, 3.3, 0.5, [(name, 20, fg, True)])
        text(s, x + 0.3, 3.25, 3.3, 2.3, [(f"·  {t}", 13, fg, False, 5) for t in items])
    text(
        s,
        MX,
        5.95,
        11.9,
        0.8,
        [
            ("과금(안)  객실(카메라) 수 기준 월 구독 + 설치 · 온보딩 1회 비용", 13, INK, True, 4),
            ("단가는 영종 운영 원가(영상 저장 · AI 연산 · 회선)를 확인한 뒤 정합니다.", 11, GRAY, False),
        ],
    )

    # 18 PRODUCT REQUIREMENTS
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 4 · SaaS 확장", "다른 시설에 팔기 위해 갖출 제품 요건")
    reqs = [
        ("멀티테넌트 · 권한", "시설별 데이터 분리\n운영자 · 매니저 · 스태프 · 보호자 역할"),
        ("온보딩", "카메라 자동 검색(ONVIF)\n룸 매핑 마법사 · 설치 가이드"),
        ("배포", "클라우드 기본\n회선이 약한 시설은 현장 엣지 장비"),
        ("과금 · 계측", "카메라 · Stay 사용량 집계\n구독 결제 · 청구서"),
        ("보안 · 개인정보", "영상 열람 기록 · 보존 기간 설정\n개인정보 처리위탁 계약 양식"),
        ("운영 지원", "카메라 끊김 감지 · 원격 점검\n무중단 업데이트"),
    ]
    for i, (h, d) in enumerate(reqs):
        cx = MX + (i % 3) * 4.05
        cy = 1.95 + (i // 3) * 2.2
        hline(s, cx, cy, 3.75, INK, 0.02)
        text(s, cx, cy + 0.18, 3.7, 0.4, [(h, 16, INK, True)])
        text(s, cx, cy + 0.7, 3.7, 1.2, [(t, 13, GRAY, False, 4) for t in d.split("\n")])
    text(s, MX, 6.45, 11.9, 0.35, [("영종 운영 기간(S0)에는 단일 시설로 쓰고, 파일럿(S1)부터 멀티테넌트 구조로 전환합니다.", 11, GRAY, False)])

    # 19 SAAS ROADMAP
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 4 · SaaS 확장", "SaaS 로드맵 — 단계마다 확인하고 넘어갑니다")
    phases = [
        ("오픈~6개월", "S0 · 영종 내부 운영", ["전 객실 가동", "알림 규칙 · 리포트 다듬기", "운영 원가 측정"], "확인  안정 가동 · 보호자 열람률"),
        ("6~12개월", "S1 · 파일럿", ["외부 시설 2~3곳", "온보딩 도구 · 멀티테넌트", "운영 사례 정리"], "확인  파일럿 유지 · 유료 의향"),
        ("12개월~", "S2 · 정식 판매", ["요금제 확정", "설치 파트너 확보", "고객 지원 체계"], "확인  유료 시설 수 · 이탈률"),
        ("이후", "S3 · 채널 확장", ["예약 플랫폼 연동", "유치원 · 병원 패키지", "다지점 고객 확보"], "확인  파트너 유입 비중"),
    ]
    hline(s, MX, 2.35, CW, INK, 0.02)
    for i, (when, name, items, gate) in enumerate(phases):
        x = MX + i * 3.05
        dot(s, x + 0.12, 2.36, 0.24, ACCENT if i == 0 else INK)
        text(s, x, 1.85, 2.8, 0.35, [(when, 13, ACCENT if i == 0 else GRAY, True)])
        text(s, x, 2.7, 2.85, 0.45, [(name, 15, INK, True)])
        text(s, x, 3.3, 2.85, 1.8, [(f"·  {t}", 13, INK, False, 6) for t in items])
        box(s, x, 5.3, 2.85, 0.6, SAND)
        text(s, x + 0.15, 5.3, 2.65, 0.6, [(gate, 11, INK, True)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX, 6.2, 11.9, 0.5, [("S0 기간에는 호텔 운영을 우선합니다. SaaS 전용 기능 개발은 S1부터 시작합니다.", 11, GRAY, False)])

    # 20 GO-TO-MARKET
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 4 · SaaS 확장", "영업 채널과 준비물")
    cols = [("채널", 2.4), ("방법", 5.6), ("준비물", 3.93)]
    rows = [
        ("직접 영업", "수도권 · 공항 권역 펫호텔 방문, 영종 견학 · 라이브 데모", "데모 계정 · 운영 사례"),
        ("설치 파트너", "CCTV 시공 업체와 제휴, 설치 시 PET-Aeye 함께 제안", "설치 가이드 · 파트너 수수료"),
        ("업계 행사", "반려동물 박람회 · 업계 단체 세미나", "부스 데모 · 소개서"),
        ("연동 제휴", "예약 · 결제 플랫폼과 연동, 앱 마켓 입점", "API 문서 · 연동 가이드"),
    ]
    table(s, MX, 1.9, cols, rows, 0.7, body=13)
    box(s, MX, 5.15, CW, 1.4, ACCENT_WASH)
    text(
        s,
        MX + 0.35,
        5.32,
        11.3,
        1.15,
        [
            ("가장 강한 영업 도구는 영종 현장입니다.", 15, INK, True, 4),
            ("실제 객실 그리드 · 보호자 리포트 · 운영 지표를 그대로 보여드립니다. 보호자 동의를 받은 화면만 사용합니다.", 12, GRAY, False),
        ],
    )

    # 21 KPI
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 5 · 관리", "핵심 지표")
    kpis = [
        ("HOTEL", "호텔 운영", ["타입별 가동률", "평균 객단가", "재이용률", "공항 유입 비중", "부가 매출 비중"]),
        ("PET-Aeye", "호텔 안 PET-Aeye", ["보호자 링크 열람률", "리포트 열람률", "안심 서비스 선택률", "알림 정확도 (오탐률)", "이벤트 대응 시간"]),
        ("SaaS", "SaaS 사업", ["활성 시설 수", "연결 카메라 수", "월 반복 매출 (MRR)", "이탈률", "온보딩 소요일"]),
    ]
    for i, (tag, name, items) in enumerate(kpis):
        x = MX + i * 4.05
        hline(s, x, 2.0, 3.75, ACCENT if i == 1 else INK, 0.03)
        text(s, x, 2.2, 3.7, 0.3, [(tag, 11, ACCENT, True)])
        text(s, x, 2.55, 3.7, 0.5, [(name, 20, INK, True)])
        for j, t in enumerate(items):
            y = 3.35 + j * 0.56
            text(s, x, y, 3.7, 0.4, [(t, 13, INK, False)])
            hline(s, x, y + 0.45, 3.75)
    text(s, MX, 6.45, 11.9, 0.35, [("목표값은 D+30 첫 리뷰에서 실제 데이터를 보고 정합니다.", 11, GRAY, False)])

    # 22 RISKS
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 5 · 관리", "리스크와 대응")
    cols = [("리스크", 3.9), ("영향", 2.4), ("대응", 5.63)]
    rows = [
        (f"{total_rooms}실 초기 가동률 저조", "고정비 부담", "3F 중대형 · 공항 패키지 우선 판매, 층별 단계 운영"),
        ("알림 오탐 · 과알림", "스태프 피로 · 신뢰 하락", "검증된 규칙만 사용, 스태프 확인 후 공유"),
        ("영상 개인정보 문제", "민원 · 법적 리스크", "고지 · 동의, 배정 룸만 열람, 열람 기록 · 보존 기간"),
        ("네트워크 · 카메라 장애", "관제 공백", "끊김 알림, 대체 순찰 SOP, 현장 녹화"),
        ("SaaS 온보딩 비용 과다", "확장 속도 저하", "설치 파트너, 자동 검색 · 매핑 도구"),
        ("호텔 운영과 SaaS 개발의 자원 충돌", "둘 다 지연", "S0에는 호텔 운영 우선, SaaS 기능은 S1부터"),
    ]
    table(s, MX, 1.9, cols, rows, 0.72, body=13)

    # 23 DECISIONS
    page += 1
    s = new_slide(prs)
    frame(s, page, "PART 5 · 관리", "결정이 필요한 사항")
    cols = [("항목", 2.9), ("선택지", 7.0), ("결정 시점", 2.03)]
    rows = [
        ("안심 서비스 요금", "전 객실 기본 포함  /  1박당 부가  /  대형 · 패밀리만 포함", "D-45"),
        ("보호자 화면", "링크형 웹  /  전용 앱", "D-45"),
        ("라이브 범위", "상시 라이브  /  정해진 시간 · 하이라이트 위주", "D-45"),
        ("영상 보존 기간", "7일  /  30일  /  Stay 종료 후 일정 기간", "D-45"),
        ("SaaS 배포 방식", "클라우드 우선  /  현장 엣지 장비 우선", "S1 전"),
        ("SaaS 브랜드", "PET-Aeye 단독  /  파우코리아 공동 표기", "S1 전"),
    ]
    table(s, MX, 1.9, cols, rows, 0.72, body=13)

    # 24 REFERENCES
    page += 1
    s = new_slide(prs)
    frame(s, page, "APPENDIX", "참고 자료")
    left = [
        ("시설", 13, ACCENT, True, 6),
        ("원 사업계획서(구칭 리조이스 · PETEL 영종) — 객실 구성 · 층별 평면 · 2차 계획 · 마케팅 계획", 12, INK, False, 14),
        ("입지 · 채널", 13, ACCENT, True, 6),
        ("인천국제공항공사 보도자료 2024.3 · 2024.11 (호텔링 · 펫라운지)", 12, INK, False, 4),
        ("인천공항 영문 안내 — 호텔링 약 13kg 이하", 12, INK, False, 4),
        ("인스파이어 리조트 이용약관 · FAQ — 반려동물 입실 불가", 12, INK, False, 4),
    ]
    right = [
        ("시장 · 수요", 13, ACCENT, True, 6),
        ("KB금융지주 경영연구소 — 펫테크 기기 사용 이유", 12, INK, False, 4),
        ("Grand View Horizon — South Korea Pet Boarding Services", 12, INK, False, 14),
        ("내부 문서", 13, ACCENT, True, 6),
        ("펫호텔_기능맵_초안.md — PET-Aeye 기능 · 화면 · API 상세", 12, INK, False, 4),
        ("프로젝트_범위_확장_초안.md — 전체 사업 범위", 12, INK, False, 4),
        ("IR_PAWKOREA_pet_hotel_yeongjong_v0.4.pptx — 투자 설명용 자료", 12, INK, False, 14),
        ("요금 · 체중 기준 · 일정은 가정입니다. 현장 실측과 운영 데이터로 갱신합니다.", 10, GRAY, False),
    ]
    text(s, MX, 2.0, 5.8, 4.6, left)
    text(s, 6.9, 2.0, 5.8, 4.6, right)

    prs.save(OUT)
    print(f"saved: {OUT}  slides={len(prs.slides)}  rooms={total_rooms}  adr={adr:,.0f}  revenues={[round(r, 1) for r in revenues]}")


if __name__ == "__main__":
    build()
