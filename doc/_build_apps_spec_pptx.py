# -*- coding: utf-8 -*-
"""PET-Aeye 앱 기획서 — 관리자 웹 · 스태프 앱 · 보호자 앱 기능 명세 deck.
디자인 헬퍼는 _build_master_pptx 에서 가져옵니다. 버전은 SPEC_VERSION 에서 관리합니다."""
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from _build_master_pptx import (
    ACCENT, ACCENT_DIM, ACCENT_WASH, ADMIN_IMG, CW, GRAD_IMG, GRAY, GUARD_IMG, H, INK, MX, NIGHT_GRAY,
    RULE, SAND, STAFF_IMG, W, WHITE, I, box, dot, hline, new_slide, section_divider, table, text,
)

SPEC_VERSION = "0.3"
DOC_DATE = "2026.10"
OUT = rf"D:\dev_pawkorea\pet-teamjang\doc\PLAN_PET-Aeye_apps_spec_v{SPEC_VERSION}.pptx"
FOOTER = "PET-Aeye 앱 기획서  ·  기능 명세 · NPU ODM · BLE 추적"
ARCH_IMG = r"D:\dev_pawkorea\pet-teamjang\doc\_mockups\out_arch_system.png"
BLE_ASSETS = r"D:\dev_pawkorea\pet-teamjang\assets\ble_tracking"

SPEC_COLS = [("ID", 0.85), ("기능", 2.35), ("상세 명세", 6.75), ("우선", 0.8), ("비고", 1.18)]

ADMIN_SPECS = [
    ("A-01", "관제 홈 · 룸 그리드", "전 객실 타일 라이브(4×2 / 4×4 전환) · 층 · 객실 타입 · 상태 필터", "P0", "데스크톱 · 월"),
    ("A-02", "AI 검출 오버레이", "개체 박스 · 행동(lying/walking/sleeping) · 신뢰도 표시, 표시 on/off", "P0", ""),
    ("A-03", "상단 KPI 바", "전체 객실 · 재실 · 바이탈 채널 · 카메라 정상 · 확인 필요 카운트", "P0", ""),
    ("A-04", "실시간 이벤트 타임라인", "우선순위 정렬 · 확인 / 담당 이관 · 처리 상태 · 2초 갱신(WebSocket)", "P0", ""),
    ("A-05", "룸 상세", "라이브 · PTZ · 녹화 타임라인 · 이벤트 이력 · 개체 메모", "P0", ""),
    ("A-06", "카메라 · 센서 관리", "ONVIF 자동 검색 · 룸 매핑 · 온라인 상태 · 오프라인 알림", "P0", "설치 시"),
    ("A-07", "바이탈 채널", "호흡 · 미동 · 재실 추세 그래프 · 개체 기준 대비 편차 표시", "P1", "웰니스"),
    ("A-08", "Stay 관리", "체크인/아웃 · 룸 · 카메라 배정 · 건강 문진 · 동의서 상태", "P1", ""),
]
ADMIN_SPECS_2 = [
    ("A-09", "보호자 안심 링크", "발송 · 회수 · 자동 만료(체크아웃) · 열람 현황", "P1", "기간 한정"),
    ("A-10", "리포트 검수 · 발송", "일일 CareReport 검수 · 발송 / Stay 졸업 리포트 생성 · 승인", "P1", ""),
    ("A-11", "알림 규칙 설정", "이벤트별 임계값 · 구역(Zone) · 야간 정책 · 에스컬레이션 대상", "P1", ""),
    ("A-12", "인력 · 교대", "층별 담당 · 근무표 · 미처리 알림 인수인계", "P1", ""),
    ("A-13", "권한 · 감사 로그", "역할(운영자/매니저/스태프) · 영상 열람 기록 · 내보내기", "P1", "보안"),
    ("A-14", "운영 통계", "가동률 · 이벤트 대응 시간 · 오탐률 · 리포트 열람률", "P2", ""),
    ("A-15", "다지점 · 테넌트", "시설별 데이터 분리 · 통합 관제 · 과금 계측", "P2", "SaaS"),
]

STAFF_SPECS = [
    ("S-01", "로그인 · 근무 시작", "교대 선택 · 담당 층 · 구역 지정 · 근무 중 상태 표시", "P0", ""),
    ("S-02", "알림 홈", "확인 필요 알림 최상단 · 푸시 · 미처리 배지 · 우선순위 색상", "P0", ""),
    ("S-03", "알림 상세", "라이브 + 오버레이 · 바이탈 바 · 개체 이력(분리불안 등) 메모", "P0", ""),
    ("S-04", "처리 액션", "현장 확인 / 정상 처리 / 이관 · 사유 선택 · 처리 시각 자동 기록", "P0", "1-탭"),
    ("S-05", "오늘 할 일", "급식 · 산책 · 투약 · 체크인 태스크 · 완료 체크 · 지연 표시", "P0", ""),
    ("S-06", "객실 리스트", "담당 구역 객실 상태 · 재실 · 알림 여부 · 라이브 진입", "P0", ""),
    ("S-07", "케어 로그", "급식량 · 배변 · 산책 · 투약 · 특이사항 · 사진 첨부", "P0", ""),
]
STAFF_SPECS_2 = [
    ("S-08", "바이탈 확인 메모", "호흡 수치 확인 · 코멘트 · 보호자 공유 전 검수 표시", "P1", "웰니스"),
    ("S-09", "보호자 공유 요청", "사진 · 메모 · 클립을 CareReport 반영 요청 → 매니저 승인", "P1", ""),
    ("S-10", "체크인 보조", "카메라 · 센서 연결 확인 · 룸 배정 확인 · 첫 사진", "P1", ""),
    ("S-11", "교대 인수인계", "미처리 알림 · 특이 개체 노트 · 다음 교대 자동 전달", "P1", ""),
    ("S-12", "공지 · 설정", "시설 공지 · 알림음 · 방해 금지 시간", "P2", ""),
]

GUARD_SPECS = [
    ("G-01", "진입 · 인증", "기간 한정 안심 링크(설치 없이) 또는 앱 로그인 · 휴대폰 인증", "P0", ""),
    ("G-02", "실시간", "배정 객실 라이브만 · 상태 헤드라인(\"편안히 쉬고 있어요\")", "P0", "배정 룸만"),
    ("G-03", "상태 요약", "호흡(웰니스) · 산책 횟수 · 컨디션 필", "P1", "비의료"),
    ("G-04", "일일 CareReport", "급식 · 산책 · 휴식 · 스태프 메모 · 사진 · 도착 알림", "P0", "스태프 검수"),
    ("G-05", "하이라이트", "일일 자동 클립 · 재생 · 저장 · 공유", "P1", ""),
    ("G-06", "Stay 졸업 리포트", "퇴실 후 종합 행동분석 · 적응 곡선 · 하이라이트 묶음", "P1", ""),
    ("G-07", "알림", "리포트 도착 · 체크아웃 안내 · (선택) 확인된 이상 공유", "P1", "확인 후만"),
]
GUARD_SPECS_2 = [
    ("G-08", "예약 · Stay 정보", "일정 · 객실 · 옵션 · 픽업/샌딩 · 안심 서비스 상태", "P1", "외부 예약 연동"),
    ("G-09", "반려견 프로필", "기본 정보 · 건강 문진 · 주의사항 · 이전 Stay 이력", "P1", ""),
    ("G-10", "개인정보 · 동의", "영상 열람 동의 · 링크 만료 안내 · 데이터 보관 기간", "P0", ""),
    ("G-11", "문의", "스태프 문의 · 카카오 채널 연결", "P2", ""),
    ("G-12", "멤버십 · 결제", "구독 · 포인트 · 결제 이력(멤버십 OS 연계)", "P2", "확장 레이어"),
]

# 메뉴 구성 — (메뉴, 하위 화면, 연결 기능 ID, 권한)
ADMIN_MENU = [
    ("관제", "룸 그리드 · 월 모드 · 층 / 상태 필터", ["A-01", "A-02", "A-03"], "전체"),
    ("이벤트", "실시간 큐 · 처리 이력 · 에스컬레이션 현황", ["A-04"], "전체"),
    ("객실", "객실 목록 · 룸 상세(라이브 · PTZ · 녹화 · 바이탈)", ["A-05", "A-07"], "전체"),
    ("Stay · 보호자", "체크인 / 아웃 · 문진 · 동의서 · 안심 링크", ["A-08", "A-09"], "운영자 · 매니저"),
    ("리포트", "일일 검수 큐 · 졸업 리포트 승인 · 발송 이력", ["A-10"], "운영자 · 매니저"),
    ("스태프", "근무표 · 층별 담당 · 인수인계", ["A-12"], "운영자 · 매니저"),
    ("장비", "카메라 · 센서 · 룸 매핑 · 온라인 상태", ["A-06"], "운영자"),
    ("통계", "가동률 · 대응 시간 · 오탐률 · 열람률", ["A-14"], "운영자 · 매니저"),
    ("설정", "알림 규칙 · 권한 · 감사 로그 · 지점(테넌트)", ["A-11", "A-13", "A-15"], "운영자"),
]
# (탭 이름, 종류: tab | util | entry, 하위 화면, 연결 기능 ID)
STAFF_MENU = [
    ("알림", "tab", ["확인 필요 큐", "알림 상세 · 라이브", "처리 액션 · 사유", "바이탈 확인 메모", "오늘 할 일 요약 3건"], ["S-02", "S-03", "S-04", "S-08"]),
    ("객실", "tab", ["담당 구역 객실", "객실 라이브", "체크인 보조", "QR 객실 스캔"], ["S-06", "S-10"]),
    ("케어 로그", "tab", ["오늘 할 일 전체", "로그 작성 · 사진", "보호자 공유 요청"], ["S-05", "S-07", "S-09"]),
    ("더보기", "tab", ["근무 시작 · 교대", "인수인계 노트", "공지", "알림 · 앱 설정"], ["S-01", "S-11", "S-12"]),
]
GUARD_MENU = [
    ("실시간", "tab", ["배정 객실 라이브", "상태 헤드라인", "상태 요약 필"], ["G-02", "G-03"]),
    ("리포트", "tab", ["일일 CareReport", "졸업 리포트(퇴실 후)", "알림함"], ["G-04", "G-06", "G-07"]),
    ("하이라이트", "tab", ["오늘의 클립", "저장 · 공유"], ["G-05"]),
    ("예약", "tab", ["Stay 정보 · 옵션", "픽업 / 샌딩", "문의 · 카카오"], ["G-08", "G-11"]),
    ("내 정보", "util", ["반려견 프로필", "동의 · 개인정보", "멤버십 · 결제", "앱 설정"], ["G-09", "G-10", "G-12"]),
]
GUARD_ENTRY = ["G-01"]

# 부가 기능 — 공통(C) · 표면별(A-16~ · S-13~ · G-13~)
COMMON_EXTRA = [
    ("C-01", "로그인 · 계정", "관리자 ID + MFA · 스태프 사번 + PIN(공용 단말 빠른 전환) · 보호자 휴대폰 인증 · 재설정", "P0", "관 · 스 · 보"),
    ("C-02", "세션 · 자동 잠금", "관리자 15분 비활성 잠금(월 모드는 열람 전용 예외) · 교대 종료 시 로그아웃", "P0", "관 · 스"),
    ("C-03", "통합 검색", "객실 번호 · 반려견 이름 · 보호자 이름 / 연락처 뒤 4자리 · 이벤트", "P0", "관 · 스"),
    ("C-04", "딥링크", "푸시 → 알림 상세 · 안심 링크 → 실시간 · 리포트 알림 → 리포트", "P0", "전체"),
    ("C-05", "연결 · 오류 · 빈 상태", "카메라 끊김 대체 화면 · 재연결 표시 · 빈 목록 안내 문구", "P0", "전체"),
    ("C-06", "약관 · 개인정보", "이용약관 · 처리방침 · 동의 철회 · 탈퇴 · 약관 버전 이력", "P0", "보 · 스"),
    ("C-07", "알림 설정", "카테고리별 on/off · 방해 금지 시간 · 알림음 · 진동", "P1", "스 · 보"),
    ("C-08", "오프라인 임시 저장", "케어 로그 · 처리 기록 로컬 큐 → 재연결 시 자동 동기화", "P1", "스"),
]
COMMON_EXTRA_2 = [
    ("C-09", "공지 · 업데이트", "시설 공지 · 강제 / 권장 업데이트 · 점검 안내", "P1", "전체"),
    ("C-10", "영상 보호", "열람자 · 시각 워터마크 · 다운로드 제한 · 캡처 시 안내", "P1", "관 · 보"),
    ("C-11", "화면 모드 · 접근성", "관제 다크 기본 · 글자 크기 · 상태는 색 + 아이콘 병기(색약 대응)", "P1", "전체"),
    ("C-12", "다국어", "한 · 영 우선 → 중 · 일 (대형 리조트 외국인 고객)", "P2", "보"),
    ("C-13", "도움말 · FAQ", "화면별 도움말 · FAQ · 고객센터 연결", "P2", "스 · 보"),
    ("C-14", "PDF · 내보내기 출력", "졸업 리포트 PDF · Stay 요약 출력 · 통계 CSV", "P2", "관 · 보"),
]
OPS_EXTRA = [
    ("A-16", "월 모드", "키오스크 전체 화면 · 층 자동 순환 · 이상 발생 타일 자동 확대", "P0", "관제실"),
    ("A-17", "경고음 · 데스크톱 알림", "긴급 이벤트 경고음 · 브라우저 알림 · 음소거 기록", "P0", ""),
    ("A-18", "단축키 · 다중 모니터", "그리드 전환 · 다음 이벤트 · 처리 단축키 · 창 분리", "P1", ""),
    ("A-19", "일괄 작업", "안심 링크 일괄 발송 · 리포트 일괄 검수 · 객실 일괄 배정", "P1", ""),
    ("A-20", "데이터 내보내기", "이벤트 · 케어 로그 · 감사 로그 CSV (권한 · 기록 남김)", "P1", "보안"),
    ("S-13", "QR 객실 스캔", "객실 문 QR → 룸 진입 · 체크인 · 케어 로그 대상 자동 선택", "P1", ""),
    ("S-14", "앱 내 촬영", "앱 안에서만 촬영 · 갤러리 미저장 · Stay 자동 태그", "P0", "보안"),
    ("S-15", "빠른 입력", "메모 템플릿 · 급식량 프리셋 · 음성 메모(선택)", "P1", ""),
    ("S-16", "긴급 호출", "원탭 매니저 / 전 스태프 호출 · 층 위치 공유", "P0", ""),
    ("S-17", "공용 단말 모드", "층별 공용 기기 · PIN 전환 · 개인 데이터 미저장", "P1", ""),
]
GUARD_EXTRA = [
    ("G-13", "공동 보호자 초대", "가족 1~3명 초대 · 열람 권한만 · 안심 링크와 같은 만료", "P1", ""),
    ("G-14", "사진 · 클립 저장 · 공유", "앨범 저장 · 카카오 공유(워터마크) · 다른 객실이 보이는 컷 제외", "P1", ""),
    ("G-15", "카카오 알림톡", "앱 미설치 보호자에게 리포트 도착 · 체크아웃 안내", "P0", "링크 모드"),
    ("G-16", "만족도 · 리뷰", "체크아웃 후 설문 · 스태프 칭찬 · NPS", "P2", ""),
    ("G-17", "재예약 · 연계 제안", "졸업 리포트에서 재예약 · 유치원 · 훈련 프로그램 제안", "P2", "확장 레이어"),
]

CORE_SPECS = ADMIN_SPECS + ADMIN_SPECS_2 + STAFF_SPECS + STAFF_SPECS_2 + GUARD_SPECS + GUARD_SPECS_2
EXTRA_SPECS = COMMON_EXTRA + COMMON_EXTRA_2 + OPS_EXTRA + GUARD_EXTRA


def check_menu_coverage():
    """핵심 기능이 모두 메뉴(또는 진입 흐름)에 연결되어 있는지 확인합니다."""
    linked = {i for m in ADMIN_MENU for i in m[2]}
    linked |= {i for m in STAFF_MENU + GUARD_MENU for i in m[3]}
    linked |= set(GUARD_ENTRY)
    core = {r[0] for r in CORE_SPECS}
    missing, unknown = sorted(core - linked), sorted(linked - core)
    assert not missing and not unknown, f"menu coverage: missing={missing} unknown={unknown}"
    ids = [r[0] for r in CORE_SPECS + EXTRA_SPECS]
    assert len(ids) == len(set(ids)), "duplicate feature id"
    return len(core)


def frame(slide, page, section, title, dark=False):
    fg = WHITE if dark else INK
    sub = NIGHT_GRAY if dark else GRAY
    text(slide, MX, 0.45, 9.5, 0.3, [(section, 11, ACCENT, True)])
    text(slide, MX, 0.8, 11.9, 1.0, [(title, 25, fg, True)])
    text(slide, MX, 7.0, 9, 0.25, [(FOOTER, 9, sub, False)])
    text(slide, 11.63, 7.0, 1.0, 0.25, [(f"{page:02d}", 9, sub, True)], PP_ALIGN.RIGHT)


def spec_slide(prs, page, section, title, rows, note=None, row_h=0.55, body=11):
    s = new_slide(prs)
    frame(s, page, section, title)
    table(s, MX, 1.75, SPEC_COLS, rows, row_h, body=body)
    if note:
        text(s, MX, 6.72, CW, 0.25, [(note, 10, GRAY, False)])
    return s


def menu_cols_slide(prs, page, section, title, menus, note):
    """모바일 앱 메뉴 구성 — 탭(진한 헤더) · 보조 메뉴(연한 헤더)를 열로 표시합니다."""
    s = new_slide(prs)
    frame(s, page, section, title)
    n, gap = len(menus), 0.2
    cw = (CW - gap * (n - 1)) / n
    for i, (name, kind, items, ids) in enumerate(menus):
        x = MX + i * (cw + gap)
        tab = kind == "tab"
        box(s, x, 1.8, cw, 0.8, INK if tab else SAND)
        text(s, x + 0.2, 1.88, cw - 0.3, 0.3, [(f"탭 {i + 1}" if tab else "우상단 메뉴", 10, ACCENT, True)])
        text(s, x + 0.2, 2.13, cw - 0.3, 0.4, [(name, 16, WHITE if tab else INK, True)])
        for j, it in enumerate(items):
            y = 2.85 + j * 0.52
            text(s, x + 0.05, y, cw - 0.1, 0.45, [(it, 12, INK, False)])
            hline(s, x, y + 0.42, cw)
        text(s, x + 0.05, 5.6, cw - 0.1, 0.3, [("연결 기능", 10, GRAY, True)])
        text(s, x + 0.05, 5.9, cw - 0.1, 0.4, [(" · ".join(ids), 12, ACCENT, True)])
    text(s, MX, 6.5, CW, 0.4, [(note, 11, GRAY, False)])
    return s


def mock_slide(prs, page, section, title, img, screens, phone=True):
    s = new_slide(prs)
    frame(s, page, section, title)
    if phone:
        s.shapes.add_picture(img, I(1.2), I(1.65), height=I(5.25))
        x, w = 5.1, 7.5
    else:
        s.shapes.add_picture(img, I(MX), I(1.8), width=I(7.6))
        x, w = 8.6, 4.0
    text(s, x, 1.85, w, 0.35, [("화면 구성", 13, ACCENT, True)])
    for i, (h, d) in enumerate(screens):
        y = 2.35 + i * 0.8
        text(s, x, y, w, 0.7, [(h, 14, INK, True, 2), (d, 11, GRAY, False)])
        hline(s, x, y + 0.7, w)
    text(s, x if phone else MX, 6.75, w if phone else CW, 0.3, [("* 화면은 구성 예시(목업)입니다.", 10, GRAY, False)])
    return s


def build():
    n_core = check_menu_coverage()
    n_extra = len(EXTRA_SPECS)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    page = 0

    # Cover
    page += 1
    s = new_slide(prs)
    box(s, 0, 0, 13.333, 0.12, ACCENT)
    text(s, MX, 0.7, 10, 0.3, [(f"APP SPEC  ·  {DOC_DATE}  ·  v{SPEC_VERSION}  ·  내부 논의용", 11, GRAY, True)])
    text(s, MX, 1.5, 11.8, 2.6, [("PET-Aeye 앱 기획서", 44, INK, True, 0), ("관리자 · 스태프 · 보호자 기능 명세", 36, ACCENT, True, 0)])
    text(s, MX, 4.15, 11.5, 0.9, [("한 데이터, 세 개의 화면 — 관제는 빠르게, 현장은 한 번에, 보호자는 안심하게.", 16, GRAY, False)])
    hline(s, MX, 5.35, CW, INK, 0.02)
    facts = [
        ("3개 표면", "관리자 웹 · 스태프 앱 · 보호자 앱"),
        (f"{n_core} + {n_extra}개 기능", "핵심 · 부가 · P0~P2 우선순위"),
        ("NPU 카메라 ODM", "바이탈형 1,000대 · Wi-Fi+PoE"),
        ("넓은 공간 BLE", "MWB01 + MKGW1-BW"),
    ]
    for i, (big, small) in enumerate(facts):
        x = MX + i * 3.0
        text(s, x, 5.6, 2.9, 0.6, [(big, 22, ACCENT if i == 0 else INK, True)])
        text(s, x, 6.2, 2.9, 0.5, [(small, 11, GRAY, False)])

    # Contents
    page += 1
    s = new_slide(prs)
    frame(s, page, "CONTENTS", "목차")
    toc = [
        ("01", "공통", "사용자 · 권한 · 시스템 구성 · 구성도 · 도메인 객체"),
        ("02", "관리자 웹", "메뉴 구성 · 관제 화면 · 기능 명세 A-01 ~ A-15"),
        ("03", "스태프 앱", "메뉴 구성 · 현장 화면 · 기능 명세 S-01 ~ S-12 · 알림 처리 플로우"),
        ("04", "보호자 앱", "메뉴 구성 · 안심 화면 · 기능 명세 G-01 ~ G-12 · 졸업 리포트 · 공개 정책"),
        ("05", "부가 기능", "공통 C-01 ~ C-14 · 표면별 A-16~ · S-13~ · G-13~ · 메뉴 점검 결과"),
        ("06", "공통 규칙", "이벤트 · 알림 정의 · 비기능 요구사항"),
        ("07", "릴리스", "MVP → P1 → P2 · 수용 기준 · KPI"),
        ("08", "NPU 카메라 ODM", "바이탈형 · 브라켓 전제 웨지 · 역할 · 비용 · 일정"),
        ("09", "넓은 공간 BLE", "MWB01 목줄 태그 · MKGW1-BW · 배치 · 체크인"),
    ]
    for i, (n, h, d) in enumerate(toc):
        y = 1.65 + i * 0.54
        text(s, MX, y, 0.8, 0.5, [(n, 18, ACCENT, True)])
        text(s, MX + 1.0, y + 0.02, 3.0, 0.5, [(h, 14, INK, True)])
        text(s, MX + 4.0, y + 0.06, 7.9, 0.5, [(d, 12, GRAY, False)])
        hline(s, MX, y + 0.48, CW)

    # ===== 01 공통 =====
    page += 1
    section_divider(prs, "01 · 공통", "한 데이터, 세 개의 화면",
                    ["사용자 역할과 권한", "카메라 · 바이탈 → AI → 이벤트 → 세 화면", "공통 도메인 객체"])

    # Roles matrix
    page += 1
    s = new_slide(prs)
    frame(s, page, "01 · 공통", "사용자 · 역할 · 권한")
    cols = [("권한", 3.2), ("운영자", 1.75), ("매니저", 1.75), ("스태프", 1.75), ("보호자", 1.75), ("비고", 1.73)]
    rows = [
        ("전 객실 라이브", "●", "●", "담당 구역", "—", ""),
        ("배정 객실 라이브", "●", "●", "●", "●", "기간 한정"),
        ("이벤트 처리", "●", "●", "●", "—", ""),
        ("리포트 검수 · 발송", "●", "●", "요청", "열람", ""),
        ("바이탈 수치", "●", "●", "●", "확인 후", "웰니스"),
        ("알림 규칙 · 설정", "●", "일부", "—", "—", ""),
        ("권한 · 감사 로그", "●", "열람", "—", "—", "보안"),
    ]
    table(s, MX, 1.85, cols, rows, 0.6, body=12)
    text(s, MX, 6.5, CW, 0.3, [("보호자에게는 스태프가 확인한 내용만, 배정된 객실만 보여줍니다.", 11, GRAY, False)])

    # System flow
    page += 1
    s = new_slide(prs)
    frame(s, page, "01 · 공통", "시스템 구성 — 수집 → 판단 → 행동 → 공유")
    flow = [
        ("수집", "NPU 카메라(ODM)\nWi-Fi+PoE · FMCW\n이벤트·클립 업로드", SAND),
        ("판단", "기기 내 검출·추적\n서버 이벤트 룰\n바이탈(웰니스)", SAND),
        ("행동", "관리자 관제 웹\n스태프 앱 알림\n처리 · 케어 로그", ACCENT_WASH),
        ("공유", "CareReport 검수\n보호자 안심 링크\n졸업 리포트", ACCENT_WASH),
    ]
    for i, (h, d, c) in enumerate(flow):
        x = MX + i * 3.02
        box(s, x, 2.0, 2.8, 2.9, c)
        text(s, x + 0.25, 2.2, 2.4, 0.5, [(f"{i + 1:02d}", 12, ACCENT, True, 2), (h, 20, INK, True)])
        text(s, x + 0.25, 3.25, 2.4, 1.6, [(t, 13, INK, False, 4) for t in d.split("\n")])
        if i < 3:
            text(s, x + 2.8, 3.2, 0.22, 0.5, [("›", 22, ACCENT, True)], PP_ALIGN.CENTER)
    text(s, MX, 5.25, CW, 1.2, [
        ("실시간: WebSocket 관제 · 모바일 푸시 · 백엔드: Stay · Room · Camera · Event · Report · AccessLink", 12, GRAY, False, 4),
        ("자격증명(RTSP 등)은 별도 보관 · 문서 · 채팅 노출 금지 · 영상 전송 암호화", 12, GRAY, False),
    ])

    # Architecture diagram
    page += 1
    s = new_slide(prs)
    frame(s, page, "01 · 공통", "관제 소프트웨어 시스템 구성도 (제안)")
    s.shapes.add_picture(ARCH_IMG, I(1.85), I(1.55), height=I(5.3))

    # Flutter client structure
    page += 1
    s = new_slide(prs)
    frame(s, page, "01 · 공통", "앱 클라이언트 — Flutter 한 코드로 웹 + 앱")
    cols = [("패키지", 2.2), ("내용", 6.4), ("빌드 대상", 3.33)]
    rows = [
        ("core", "OpenAPI → Dart 클라이언트 자동 생성 · 인증 · 도메인 모델 · WebSocket · 오류 화면 · 다국어", "공통"),
        ("ui_kit", "디자인 토큰 · 상태 칩 · 바이탈 바 · 라이브 플레이어(WebRTC) · 비의료 표기 컴포넌트", "공통"),
        ("staff_app", "알림 · 처리 · 케어 로그 · 객실 · 교대 (S-01 ~ S-17)", "Android 우선 · iOS · Web"),
        ("guardian_app", "실시간 · 리포트 · 하이라이트 · 예약 · 내 정보 (G-01 ~ G-17)", "Android · iOS · Web(안심 링크)"),
        ("관리자 웹", "월 모드 대량 라이브 그리드 · MVP 화면 재사용", "React 유지 (P1 이후 재판단)"),
    ]
    table(s, MX, 1.8, cols, rows, 0.62, body=12)
    text(s, MX, 5.55, CW, 1.2, [
        ("Riverpod 상태 관리 · go_router 딥링크 · flutter_webrtc 라이브 · FCM 푸시 · 로컬 DB 오프라인 저장", 12, INK, False, 6),
        ("안심 링크(Flutter Web)는 첫 화면을 가볍게 하고 리포트 · 하이라이트는 지연 로딩 — 목표: 첫 화면 3초 이내(LTE)", 12, GRAY, False),
    ])

    # Domain objects
    page += 1
    s = new_slide(prs)
    frame(s, page, "01 · 공통", "공통 도메인 객체")
    cols = [("객체", 2.3), ("주요 필드", 6.0), ("사용 화면", 3.63)]
    rows = [
        ("Guardian · Pet", "보호자 연락처 · 반려견 프로필 · 체중 · 건강 문진 · 주의사항", "관리자 · 보호자"),
        ("Stay", "체크인/아웃 · 객실 · 옵션 · 상태 · 동의서", "전 화면"),
        ("Room · Camera · Sensor", "층 · 타입 · 카메라 매핑 · 센서 · 온라인 상태", "관리자 · 스태프"),
        ("PetEvent", "타입 · 시각 · 신뢰도 · 클립 · 처리 상태 · 담당", "관리자 · 스태프"),
        ("CareTask · CareLog", "급식 · 산책 · 투약 · 배변 · 메모 · 사진", "스태프 → 리포트"),
        ("VitalSample", "호흡 · 미동 · 재실 · 기준값 · 편차 (웰니스)", "관리자 · 스태프 · 보호자"),
        ("CareReport · GradReport", "일일 요약 · 검수자 · 발송 / 종합 행동분석", "보호자"),
        ("AccessLink", "토큰 · 대상 Stay · 만료 · 열람 로그", "보호자"),
    ]
    table(s, MX, 1.8, cols, rows, 0.55, body=11)

    # ===== 02 관리자 웹 =====
    page += 1
    section_divider(prs, "02 · 관리자 웹", "112실을 한 화면에서",
                    ["룸 그리드 관제 · AI 오버레이 · 바이탈 칩", "이벤트 타임라인과 처리 이관", "Stay · 링크 · 리포트 · 설정 관리"])
    page += 1
    s = new_slide(prs)
    frame(s, page, "02 · 관리자 웹", "메뉴 구성 — 좌측 GNB 9개")
    cols = [("메뉴", 2.0), ("하위 화면", 5.6), ("연결 기능", 2.3), ("권한", 2.03)]
    rows = [(f"{i + 1}  {m}", sub, " · ".join(ids), role) for i, (m, sub, ids, role) in enumerate(ADMIN_MENU)]
    table(s, MX, 1.75, cols, rows, 0.48, body=11)
    text(s, MX, 6.62, CW, 0.3, [("상단 고정: 통합 검색(C-03) · 알림 벨 · 월 모드 전환(A-16) · 프로필 / 로그아웃(C-01)  ·  스태프 권한은 담당 구역만 표시", 10, GRAY, False)])
    page += 1
    mock_slide(prs, page, "02 · 관리자 웹", "관제 홈 — 화면 구성", ADMIN_IMG, [
        ("상단 KPI 바", "객실 · 재실 · 바이탈 · 카메라 · 확인 필요"),
        ("룸 그리드", "타일 라이브 · 검출 박스 · 상태 칩"),
        ("바이탈 칩", "호흡 · 재실 · 수면 (웰니스)"),
        ("이상 강조", "확인 필요 객실 테두리 강조"),
        ("이벤트 사이드바", "우선순위 · 처리 · 로그"),
    ], phone=False)
    page += 1
    spec_slide(prs, page, "02 · 관리자 웹", "기능 명세 ① — 관제 · 객실 · 센서", ADMIN_SPECS,
               "P0 = 오픈 필수 · P1 = 오픈 후 3개월 · P2 = SaaS 확장 단계")
    page += 1
    spec_slide(prs, page, "02 · 관리자 웹", "기능 명세 ② — 운영 · 보안 · 확장", ADMIN_SPECS_2)

    # ===== 03 스태프 앱 =====
    page += 1
    section_divider(prs, "03 · 스태프 앱", "알림은 위에, 처리는 한 번에",
                    ["확인 필요 알림 우선 · 라이브 + 바이탈 바", "현장 확인 / 정상 처리 / 이관", "담당 층 할 일 · 케어 로그 · 인수인계"])
    page += 1
    menu_cols_slide(prs, page, "03 · 스태프 앱", "메뉴 구성 — 하단 탭 4개", STAFF_MENU,
                    "상단 고정: 근무 상태 · 담당 층 · 긴급 호출(S-16) · 검색(C-03)  ·  알림 탭의 '오늘 할 일'은 요약 3건, 전체는 케어 로그 탭")
    page += 1
    mock_slide(prs, page, "03 · 스태프 앱", "알림 홈 — 화면 구성", STAFF_IMG, [
        ("확인 필요 알림", "가장 위 · 라이브 + 검출 박스"),
        ("바이탈 바", "호흡 수치 · 재실 (웰니스)"),
        ("개체 맥락", "분리불안 등 이력 메모"),
        ("1-탭 처리", "현장 확인 / 정상 처리"),
        ("오늘 할 일", "급식 · 산책 · 투약 · 체크인"),
    ])
    page += 1
    spec_slide(prs, page, "03 · 스태프 앱", "기능 명세 ① — 알림 · 처리 · 태스크", STAFF_SPECS)
    page += 1
    spec_slide(prs, page, "03 · 스태프 앱", "기능 명세 ② — 바이탈 · 공유 · 교대", STAFF_SPECS_2)

    # Alert handling flow
    page += 1
    s = new_slide(prs)
    frame(s, page, "03 · 스태프 앱", "알림 처리 플로우")
    steps = [("수신", "푸시 · 알림 홈 최상단"), ("판단", "라이브 · 바이탈 · 개체 메모"), ("처리", "현장 확인 / 정상 / 이관"), ("기록", "사유 · 시각 · 케어 로그"), ("공유", "매니저 승인 시 리포트 반영")]
    hline(s, MX, 2.75, CW, INK, 0.02)
    for i, (h, d) in enumerate(steps):
        x = MX + i * 2.43
        dot(s, x + 0.14, 2.76, 0.28, ACCENT if i in (0, 4) else INK)
        text(s, x, 2.05, 2.3, 0.4, [(f"STEP {i + 1}", 11, ACCENT if i in (0, 4) else GRAY, True)])
        text(s, x, 3.1, 2.3, 0.5, [(h, 18, INK, True)])
        text(s, x, 3.65, 2.25, 0.9, [(d, 12, GRAY, False)])
    box(s, MX, 4.85, 5.85, 1.6, SAND)
    text(s, MX + 0.3, 5.05, 5.3, 1.3, [
        ("처리 기준", 13, INK, True, 5),
        ("· 미처리 5분 → 매니저 에스컬레이션", 12, INK, False, 3),
        ("· 긴급(미검출 · 이탈) → 즉시 전 스태프", 12, INK, False),
    ])
    box(s, 6.75, 4.85, 5.88, 1.6, ACCENT_WASH)
    text(s, 7.05, 5.05, 5.3, 1.3, [
        ("보호자 공유 원칙", 13, ACCENT, True, 5),
        ("· 원시 알림은 보호자에게 보내지 않습니다", 12, INK, False, 3),
        ("· 스태프 확인 + 매니저 승인 후 리포트로만", 12, INK, False),
    ])

    # ===== 04 보호자 앱 =====
    page += 1
    section_divider(prs, "04 · 보호자 앱", "맡기는 순간부터, 보이는 안심을",
                    ["배정 객실 라이브 · 상태 헤드라인", "일일 CareReport · 하이라이트", "Stay 졸업 리포트 · 기간 한정 링크"])
    page += 1
    menu_cols_slide(prs, page, "04 · 보호자 앱", "메뉴 구성 — 하단 탭 4개 + 내 정보", GUARD_MENU,
                    "진입(G-01): 안심 링크(설치 없음)는 실시간 · 리포트 · 하이라이트 3개 탭만, 예약 · 내 정보는 앱 로그인 후 열립니다.")
    page += 1
    mock_slide(prs, page, "04 · 보호자 앱", "실시간 홈 — 화면 구성", GUARD_IMG, [
        ("상태 헤드라인", "\"초코는 지금 편안히 쉬고 있어요\""),
        ("배정 객실 라이브", "내 반려견 방만"),
        ("상태 요약 필", "호흡(웰니스) · 산책 · 컨디션"),
        ("오늘의 CareReport", "급식 · 산책 · 휴식 · 메모"),
        ("하이라이트 · 졸업 리포트", "클립 · 퇴실 후 종합"),
    ])
    page += 1
    spec_slide(prs, page, "04 · 보호자 앱", "기능 명세 ① — 실시간 · 리포트", GUARD_SPECS)
    page += 1
    spec_slide(prs, page, "04 · 보호자 앱", "기능 명세 ② — 예약 · 프로필 · 동의", GUARD_SPECS_2)

    # Graduation report spec
    page += 1
    s = new_slide(prs)
    frame(s, page, "04 · 보호자 앱", "G-06 Stay 졸업 리포트 — 상세 명세")
    s.shapes.add_picture(GRAD_IMG, I(1.0), I(1.65), height=I(5.25))
    cols = [("구성 요소", 2.3), ("데이터 · 생성 규칙", 4.2)]
    rows = [
        ("요약 지표", "투숙일수 · 총 활동량 · 수면 시간 · 산책 횟수"),
        ("적응 곡선", "일별 안정도 점수(첫날 대비) 막대"),
        ("행동 프로파일", "활동 성향 · 사회성 · 분리불안 신호 칩"),
        ("하이라이트", "이벤트 타임라인 상위 클립 3~5개"),
        ("케어 로그 요약", "급식 · 투약 · 특이사항"),
        ("스태프 코멘트", "담당 · 훈련사 코멘트 (검수)"),
        ("발송", "체크아웃 후 24시간 내 · 링크 30일"),
    ]
    table(s, 5.9, 1.8, cols, rows, 0.6, body=11)

    # Exposure policy
    page += 1
    s = new_slide(prs)
    frame(s, page, "04 · 보호자 앱", "안심 링크 · 공개 정책")
    cols = [("항목", 2.6), ("정책", 6.2), ("근거", 3.13)]
    rows = [
        ("열람 범위", "배정 객실 라이브만 · 다른 객실 · 공용 구역 비노출", "개인정보 · 타 고객 보호"),
        ("링크 수명", "체크인 시 발급 · 체크아웃 시 자동 만료 · 수동 회수", "기간 한정"),
        ("인증", "휴대폰 인증 · 토큰 재발급 시 기존 토큰 폐기", "도용 방지"),
        ("바이탈 노출", "웰니스 참고값 · 스태프 확인 후 · 진단 표현 금지", "비의료"),
        ("알림", "원시 이벤트 비발송 · 확인된 내용만 리포트로", "불안 · 오해 방지"),
        ("기록", "모든 열람 감사 로그 · 보관 기간 정책 고지", "보안"),
    ]
    table(s, MX, 1.85, cols, rows, 0.65, body=12)

    # ===== 05 부가 기능 =====
    page += 1
    section_divider(prs, "05 · 부가 기능", "메뉴 밖에서 받쳐 주는 기능",
                    [f"공통 부가 기능 C-01 ~ C-{len(COMMON_EXTRA + COMMON_EXTRA_2):02d}", "관리자 · 스태프 · 보호자 표면별 부가 기능", "메뉴 · 기능 점검 결과"])
    page += 1
    spec_slide(prs, page, "05 · 부가 기능", "공통 부가 기능 ① — 계정 · 검색 · 연결", COMMON_EXTRA,
               "비고 = 적용 표면 (관 = 관리자 웹 · 스 = 스태프 앱 · 보 = 보호자 앱)")
    page += 1
    spec_slide(prs, page, "05 · 부가 기능", "공통 부가 기능 ② — 공지 · 보호 · 접근성", COMMON_EXTRA_2)
    page += 1
    spec_slide(prs, page, "05 · 부가 기능", "관리자 · 스태프 부가 기능", OPS_EXTRA, row_h=0.45, body=10)
    page += 1
    spec_slide(prs, page, "05 · 부가 기능", "보호자 부가 기능", GUARD_EXTRA,
               "공유 · 저장 콘텐츠에도 '다른 객실 비노출' · 워터마크(C-10) 원칙을 동일하게 적용합니다.")

    page += 1
    s = new_slide(prs)
    frame(s, page, "05 · 부가 기능", "메뉴 · 기능 점검 결과")
    cols = [("점검 항목", 2.6), ("확인 내용", 4.9), ("조치", 4.43)]
    rows = [
        ("메뉴 ↔ 기능 연결", f"핵심 기능 {n_core}개 모두 메뉴 또는 진입 흐름에 연결 (빌드 시 자동 검사)", "누락 0건 유지"),
        ("관리자 GNB", "목업은 관제 홈만 있고 메뉴가 없음", "좌측 GNB 9개 메뉴를 목업에 추가"),
        ("보호자 내 정보", "프로필 · 동의 · 멤버십이 하단 탭에 진입점 없음", "우상단 '내 정보' 메뉴로 배치"),
        ("안심 링크 모드", "설치 없는 웹에서 예약 · 내 정보 사용 불가", "3개 탭만 노출 · 앱 설치 안내"),
        ("오늘 할 일 중복", "스태프 알림 홈 하단과 케어 로그 탭에 모두 존재", "홈은 요약 3건 · 전체는 케어 로그"),
        ("현장 안전 기능", "긴급 호출 · 경고음 · 월 모드가 기존 명세에 없음", "S-16 · A-17 · A-16 을 P0으로 추가"),
        ("앱 기본 기능", "로그인 · 약관 · 딥링크 · 오류 화면 미정의", "C-01 ~ C-06 을 P0으로 추가"),
    ]
    table(s, MX, 1.8, cols, rows, 0.6, body=11)

    # ===== 06 공통 규칙 =====
    page += 1
    section_divider(prs, "06 · 공통 규칙", "무엇을 누구에게 언제 알리나",
                    ["이벤트 · 알림 정의", "비기능 요구사항"])
    page += 1
    s = new_slide(prs)
    frame(s, page, "06 · 공통 규칙", "이벤트 · 알림 정의")
    cols = [("이벤트", 2.4), ("트리거 (초기값 · 30일 데이터로 조정)", 4.4), ("스태프", 1.6), ("관리자", 1.5), ("보호자", 2.03)]
    rows = [
        ("E1 장시간 무움직임", "N분 이상 정지 · 호흡 신호 동반 확인", "즉시", "타임라인", "확인 후"),
        ("E2 장시간 활동", "25분 이상 연속 이동 · 배회", "즉시", "타임라인", "확인 후"),
        ("E3 호흡 편차", "개체 기준 대비 편차 지속 (웰니스)", "즉시", "타임라인", "확인 후"),
        ("E4 미검출 · 이탈", "객실 내 개체 미검출", "긴급 · 전원", "강조", "비노출"),
        ("E5 장비 오프라인", "카메라 · 센서 연결 끊김", "—", "즉시", "비노출"),
        ("E6 케어 완료", "급식 · 산책 · 투약 체크", "—", "로그", "리포트"),
        ("E7 소리 · 짖음", "지속 짖음 (옵션)", "알림", "타임라인", "확인 후"),
    ]
    table(s, MX, 1.8, cols, rows, 0.6, body=11)

    page += 1
    s = new_slide(prs)
    frame(s, page, "06 · 공통 규칙", "비기능 요구사항")
    cols = [("영역", 2.3), ("요구사항 (목표값)", 6.6), ("측정", 3.03)]
    rows = [
        ("지연", "관제 라이브 2초 이내(LAN) · 알림 도달 5초 이내", "모니터링"),
        ("가용성", "관제 99.5% · 녹화 무중단 · 장비 이중화(관제실)", "월간"),
        ("정확도", "오탐률 · 미탐률 지표 관리 · 임계값 30일 튜닝", "D+30 리뷰"),
        ("보안", "전송 암호화 · 관리자 MFA · 역할 기반 권한 · 감사 로그", "점검"),
        ("개인정보", "보관 기간 정책 · 동의 기록 · 링크 만료", "점검"),
        ("확장", "멀티테넌트 데이터 분리 · 다지점 · 과금 계측", "S1~"),
        ("단말", "관리자: 데스크톱/월(React) · 스태프 · 보호자: Flutter 한 코드로 Android 우선 · iOS · Web", "안심 링크 첫 화면 3초"),
    ]
    table(s, MX, 1.8, cols, rows, 0.6, body=11)

    # ===== 07 릴리스 =====
    page += 1
    section_divider(prs, "07 · 릴리스", "오픈에 꼭 필요한 것부터",
                    ["MVP → P1 → P2 단계", "수용 기준 · KPI"])
    page += 1
    s = new_slide(prs)
    frame(s, page, "07 · 릴리스", "릴리스 로드맵")
    phases = [
        ("MVP · P0", "오픈 D-30", ["A-01~06 관제 · 센서", "S-01~07 알림 · 태스크 · 로그", "G-01·02·04·10 라이브 · 리포트 · 동의", "부가: C-01~06 · A-16·17 · S-14·16 · G-15"], ACCENT),
        ("P1", "오픈 ~ +3개월", ["바이탈 채널 · 메모 · 상태 필", "Stay · 링크 · 리포트 검수", "하이라이트 · 졸업 리포트 · 예약", "부가: 알림 설정 · 오프라인 · 공동 보호자"], INK),
        ("P2", "+6개월 ~", ["운영 통계 · 다지점 · 테넌트", "문의 · 멤버십 · 결제", "SaaS 온보딩 · 과금", "부가: 다국어 · FAQ · PDF · 리뷰"], GRAY),
    ]
    for i, (name, when, items, c) in enumerate(phases):
        x = MX + i * 4.05
        hline(s, x, 1.95, 3.75, c, 0.04)
        text(s, x, 2.15, 3.7, 0.4, [(when, 12, c, True)])
        text(s, x, 2.55, 3.7, 0.6, [(name, 24, INK, True)])
        text(s, x, 3.4, 3.7, 2.4, [("· " + it, 13, INK, False, 8) for it in items])
    text(s, MX, 6.4, CW, 0.4, [("MVP 범위는 '관제 → 처리 → 보호자 리포트' 한 바퀴가 도는 최소 단위입니다.", 12, GRAY, False)])

    page += 1
    s = new_slide(prs)
    frame(s, page, "07 · 릴리스", "수용 기준 · KPI")
    cols = [("화면", 2.0), ("수용 기준 (오픈 전)", 5.6), ("KPI (오픈 후)", 4.33)]
    rows = [
        ("관리자 웹", "전 객실 카메라 매핑 · 그리드 2초 갱신 · 이벤트 처리 이관 동작", "이벤트 대응 시간 · 오탐률"),
        ("스태프 앱", "알림 5초 도달 · 1-탭 처리 기록 · 태스크 완료 체크", "미처리 알림 수 · 태스크 완료율"),
        ("보호자 앱", "배정 객실만 노출 · 링크 자동 만료 · 리포트 검수 후 발송", "링크 열람률 · 리포트 열람률"),
        ("공통", "감사 로그 기록 · 권한 분리 · 비의료 표기 검수", "안심 옵션 선택률 · 재이용률"),
    ]
    table(s, MX, 1.85, cols, rows, 0.85, body=12)
    text(s, MX, 6.2, CW, 0.5, [("목표값은 오픈 후 D+30 리뷰에서 확정합니다. 상세 서비스 범위 · IT 툴은 별도 문서(PET-Aeye_서비스범위_IT툴_기획서_초안.md)를 따릅니다.", 11, GRAY, False)])

    # ===== 08 NPU 카메라 ODM =====
    page += 1
    section_divider(prs, "08 · NPU 카메라 ODM", "기기에서 추론하고, 앱은 의미만",
                    ["바이탈형 단일 SKU · 1차 500 + 2차 500", "브라켓 전제 · 미니 웨지 본체", "역할 · 비용 · 일정"])

    page += 1
    s = new_slide(prs)
    frame(s, page, "08 · NPU 카메라 ODM", "요청 개요 — 바이탈형 일체형 IP 카메라")
    cols = [("항목", 2.4), ("내용", 9.53)]
    rows = [
        ("제품 (가칭)", "PET-Aeye NPU Camera · Wi-Fi + 이더넷 + PoE(802.3af) 단일 제품"),
        ("칩 · 추론", "Axera AX630C(또는 동등) · 기기 내 탐지·포즈·추적 → 결과·이벤트만 서버 전송"),
        ("주력 SKU", "바이탈형 = NPU + 60GHz FMCW · JSON/MQTT API · 전량 동일 SKU"),
        ("수량", "1차 약 500대(2027 · 영종·홈) · 2차 약 500대(횡성 전) · 합 약 1,000대"),
        ("샘플", "EVT/DVT 바이탈 20~50대 · AI · OTA · 레이더·차별 기능 검증"),
        ("펌웨어", "OTA 필수 · AI 모델(.axmodel)은 발주처 제공 · 서명·롤백"),
        ("관련 RFQ", "doc/PET-Aeye_NPU카메라_ODM_견적요청서_초안.md"),
    ]
    table(s, MX, 1.75, cols, rows, 0.58, body=12)
    text(s, MX, 6.35, CW, 0.4, [("바이탈 = 웰니스 참고값(비의료). 표준형(레이더 없음)은 비교용 단가만 요청 · 발주 의무 없음.", 11, GRAY, False)])

    page += 1
    s = new_slide(prs)
    frame(s, page, "08 · NPU 카메라 ODM", "역할 분담 — ODM ↔ 발주처(PET-Aeye)")
    cols = [("영역", 2.2), ("ODM (기기)", 4.9), ("발주처 (우리)", 4.83)]
    rows = [
        ("하드웨어", "미니 웨지·브라켓·셔터·레이더 · KC", "사양·목표 단가 · BOM 승인"),
        ("AI", "NPU 런타임 · .axmodel 로드·추론", "학습 · Pulsar2 변환 · OTA 배포"),
        ("설정·제어 API", "수신·적용·상태 조회 구현", "스키마·인증 정의 · 앱/서버 중계"),
        ("텔레메트리", "비전·레이더·모드 JSON → 우리 서버만", "수집 · Stay 매핑 · 알림 문구"),
        ("OTA", "기기 에이전트 · 서명 검증·롤백", "OTA 서버 · 카나리 · 키 관리"),
        ("차별 기능", "모드·외출·블러·클립·폴백 실행", "정책·UI · 일일 카드·그래프"),
        ("앱·관제", "불필요 (기기만)", "관리자 · 스태프 · 보호자 전 화면"),
    ]
    table(s, MX, 1.75, cols, rows, 0.58, body=11)
    text(s, MX, 6.35, CW, 0.4, [("앱이 칩을 직접 호출하지 않습니다. 설정(호텔/홈·외출·블러 등)은 전부 서버 → 기기 API.", 11, GRAY, False)])

    page += 1
    s = new_slide(prs)
    frame(s, page, "08 · NPU 카메라 ODM", "브라켓 전제 — 벽 + 천장 겸용 마운트")
    cols = [("항목", 2.3), ("기준", 9.63)]
    rows = [
        ("원칙", "동일 본체 · 벽면·천장(면) 마운트 모두 지원 · 키트만 분기"),
        ("본체", "미니 웨지 · 단일 마운트 페이스 · 후면 배선 · 화면 회전(corridor)"),
        ("B-Wall", "벽 마운트 · 오픈 룸 1.4~1.7m 하향 · 룸당 1대"),
        ("B-Ceiling", "천장 면부착 · 하향 · 한 공간=한 시야일 때만"),
        ("B-Gable", "케이지 지붕 안쪽(천장형) · 칸당 1대 · 큰 구멍 X"),
        ("B-Frame / Home", "철망 클램프 · 홈은 벽 또는 천장"),
        ("운용", "브라켓 고정 · 본체만 교체 · 벽↔천장 키트 교체로 전환"),
    ]
    table(s, MX, 1.7, cols, rows, 0.58, body=12)
    text(s, MX, 6.3, CW, 0.4, [("RFQ §3.1 · 펫호텔 권장: 오픈룸=벽 · 케이지=천장형 면부착 · 본체는 하나.", 11, GRAY, False)])

    page += 1
    s = new_slide(prs)
    frame(s, page, "08 · NPU 카메라 ODM", "하드웨어 · 차별 기능 (펌웨어 기본 포함)")
    box(s, MX, 1.75, 5.9, 4.85, SAND)
    text(s, MX + 0.25, 1.9, 5.4, 0.4, [("하드웨어 (목표)", 14, ACCENT, True)])
    text(s, MX + 0.25, 2.4, 5.4, 3.9, [
        ("· 4MP · 1/1.8\"급 저조도 · 850nm IR", 12, INK, False, 5),
        ("· Wi-Fi 6 + 유선 랜 + PoE 단일 기판", 12, INK, False, 5),
        ("· 미니 웨지 + 브라켓 · 물리 셔터", 12, INK, False, 5),
        ("· 60GHz FMCW · 재실·호흡·미동", 12, INK, False, 5),
        ("· 온습도 · 촬영 LED · 셔터 상태 보고", 12, INK, False, 5),
        ("· RTSP(H.265) · 사람 블러 on/off", 12, INK, False),
    ])
    box(s, 6.85, 1.75, 5.78, 4.85, ACCENT_WASH)
    text(s, 7.1, 1.9, 5.3, 0.4, [("차별 기능 §4.5 (ODM 실행)", 14, ACCENT, True)])
    text(s, 7.1, 2.4, 5.3, 3.9, [
        ("· 호텔 / 홈 모드 프로필", 12, INK, False, 5),
        ("· 외출 모드 + 셔터 연동", 12, INK, False, 5),
        ("· 이벤트 클립 · 회선 폴백 재전송", 12, INK, False, 5),
        ("· 수면·활동 시계열 필드 보고", 12, INK, False, 5),
        ("· (가산점) 홈 모드 양방향 음성", 12, INK, False, 5),
        ("· Guard·브라켓 키트 = 분리 견적", 12, INK, False),
    ])

    page += 1
    s = new_slide(prs)
    frame(s, page, "08 · NPU 카메라 ODM", "비용 ① — 대당 원가 · 목표 단가 (내부 가정)")
    cols = [("등급", 2.0), ("구성", 4.5), ("대당 원가 (대략)", 2.6), ("비고", 2.83)]
    rows = [
        ("기본형", "NPU only", "약 4~4.5만원", "비교 · 발주 안 함"),
        ("표준형", "NPU · Wi-Fi+PoE", "약 7~7.7만원 / ~10.3만원*", "*통합형 산정용"),
        ("바이탈형 ★", "표준 + 60GHz FMCW", "약 10.5~15만원", "주력 · RFQ 목표"),
        ("RFQ 목표", "로트 500 · 누적 1,000대", "USD 85~110 / 12~15만원대", "VAT·관세 별도"),
    ]
    table(s, MX, 1.75, cols, rows, 0.62, body=12)
    text(s, MX, 4.55, CW, 2.2, [
        ("* 영종·횡성 CAPEX 표의 NPU 통합형 대당 약 10.3만원은 Wi-Fi+PoE 기준(표준형 근처). 바이탈 전량 발주 시 대당·합계는 ODM 회신으로 갱신.", 12, GRAY, False, 6),
        ("NRE(개발·금형·OTA·바이탈 API·KC·샘플)는 홈·영종·횡성이 공유 — 시설 CAPEX에서 제외하고 별도 항으로 견적.", 12, GRAY, False, 6),
        ("견적 요청: 500 / 1,000 / 3,000대 단가 · OTA 포함·미포함 차액 · NRE 항목별.", 12, INK, False),
    ])

    page += 1
    s = new_slide(prs)
    frame(s, page, "08 · NPU 카메라 ODM", "비용 ② — 시설 CAPEX (카메라 중심 · NRE 제외)")
    cols = [("시설", 2.3), ("NPU 대수 (대략)", 2.5), ("카메라·구축 합 (대략)", 3.5), ("비고", 3.63)]
    rows = [
        ("영종 펫호텔", "실내 150 · 실외 시중 6", "약 0.80억 (v0.4)", "Wi-Fi 연결 · 관제 최소"),
        ("횡성 펫 리조트", "실내 310 · 실외 30", "카메라 구성 문서 기준", "PoE · IT 거점·GPU 별도"),
        ("홈서비스", "구독 대수에 따름", "대당 원가 × 대수", "클라우드 OTA"),
        ("ODM 1·2차", "합 약 1,000대 바이탈", "대당 × 1,000 + NRE", "동일 SKU 분할 납품"),
    ]
    table(s, MX, 1.75, cols, rows, 0.65, body=12)
    text(s, MX, 4.7, CW, 2.0, [
        ("근거: PET-Aeye_CCTV_서버_비용산정_초안.md · 횡성펫리조트_카메라_구성_초안.md · 홈서비스_기획_초안.md", 11, GRAY, False, 6),
        ("횡성 HQ IT(서버실·관제·GPU·사내 LLM 등)는 카메라 CAPEX와 별도(개관 시 약 4.3~6.0억 구간 · 문서 참고).", 12, INK, False, 6),
        ("마스터 '약 1.7억' 서버형 숫자는 NPU 양산 확정 후 본 기준으로 교체 예정.", 11, GRAY, False),
    ])

    page += 1
    s = new_slide(prs)
    frame(s, page, "08 · NPU 카메라 ODM", "일정 · 견적 회신 요청 항목")
    cols = [("마일스톤", 3.0), ("목표", 4.0), ("산출물", 4.93)]
    rows = [
        ("EVT 샘플", "+8~12주 (계약 후)", "바이탈 동작 · 영상·추론 · 레이더"),
        ("DVT + 시연", "+4~8주", "OTA · API · 모드·블러·폴백 · 온습도"),
        ("1차 양산", "2027", "바이탈 약 500대 · KC(레이더 포함)"),
        ("2차 양산", "횡성 준공 전", "바이탈 약 500대 · 동일 SKU"),
    ]
    table(s, MX, 1.75, cols, rows, 0.55, body=12)
    text(s, MX, 4.3, CW, 2.4, [
        ("회신에 반드시: 회사·실적 · 사양 대비표 · §6 견적 · MOQ·결제 · 일정 · OTA 차액 · 바이탈 모듈·API · 차별 기능 · 설정 API · Longevity", 12, INK, False, 6),
        ("발송 전: 법인·담당 · NDA · 목표 단가 대외 공개 범위 확정. 본 RFQ는 발주 확정 아님.", 12, GRAY, False),
    ])

    # ===== 09 넓은 공간 BLE =====
    page += 1
    section_divider(prs, "09 · 넓은 공간 BLE", "카메라는 행동, BLE는 개체 ID",
                    ["Minew MWB01 목줄 태그", "MOKO MKGW1-BW 게이트웨이", "플레이그라운드 · 루프탑"])

    page += 1
    s = new_slide(prs)
    frame(s, page, "09 · 넓은 공간 BLE", "왜 · 무엇 — 넓은 야외 ID 유지")
    cols = [("항목", 2.2), ("내용", 9.73)]
    rows = [
        ("문제", "지상 플레이그라운드 · 옥상 루프탑은 원거리·혼잡으로 비전만으로 개체 ID 유지가 어려움"),
        ("태그", "Minew MWB01 · iBeacon · adv 300ms · UUID/Major/Minor · MAC allowlist"),
        ("게이트웨이", "MOKO MKGW1-BW · MQTT/TLS · PoE · IP20→옥외 IP65 함체 필수"),
        ("착용", "목줄 관통·몸쪽 밀착 실리콘 케이스(매달림·키링 금지)"),
        ("측위", "L1 Zone + L2 Gate (cm급·UWB 아님) · 이탈 알림 ≤15s 목표"),
        ("문서", "PET-Aeye_넓은공간_BLE개체추적_초안.md v0.3 · assets/ble_tracking/"),
    ]
    table(s, MX, 1.7, cols, rows, 0.58, body=12)
    text(s, MX, 5.7, CW, 1.0, [
        ("앱: bind/unbind · zone.enter/exit · geofence.breach · gw.offline · 비전 조인(stay_id).", 12, INK, False, 4),
        ("구역당 BOM 약 496만원 가정(태그·케이스·GW·함체·PoE·데스크). 카메라 CAPEX 별도.", 11, GRAY, False),
    ])

    page += 1
    s = new_slide(prs)
    frame(s, page, "09 · 넓은 공간 BLE", "BOM · KPI · 이벤트 (v0.3)")
    box(s, MX, 1.7, 3.9, 4.9, SAND)
    text(s, MX + 0.2, 1.85, 3.5, 0.35, [("구역당 BOM", 14, ACCENT, True)])
    text(s, MX + 0.2, 2.35, 3.5, 4.0, [
        ("· MWB01 ×80", 12, INK, False, 4),
        ("· 관통 케이스 ×120", 12, INK, False, 4),
        ("· MKGW1-BW ×5", 12, INK, False, 4),
        ("· IP65 함체·폴·PoE", 12, INK, False, 4),
        ("· 데스크 키트 · 목줄", 12, INK, False, 4),
        ("· 합 ≈ 496만원", 12, ACCENT, True),
    ])
    box(s, 4.85, 1.7, 3.9, 4.9, ACCENT_WASH)
    text(s, 5.05, 1.85, 3.5, 0.35, [("KPI (MVP)", 14, ACCENT, True)])
    text(s, 5.05, 2.35, 3.5, 4.0, [
        ("· Zone L1 + Gate L2", 12, INK, False, 4),
        ("· 게이트 검출 ≥95%", 12, INK, False, 4),
        ("· 이탈 알림 ≤15s", 12, INK, False, 4),
        ("· 동시 태그 ≥80", 12, INK, False, 4),
        ("· GW 가용 ≥99%", 12, INK, False, 4),
        ("· UWB/AoA = Phase2", 12, GRAY, False),
    ])
    box(s, 9.0, 1.7, 3.63, 4.9, SAND)
    text(s, 9.2, 1.85, 3.25, 0.35, [("도메인 이벤트", 14, ACCENT, True)])
    text(s, 9.2, 2.35, 3.25, 4.0, [
        ("· ble.tag.bind/unbind", 11, INK, False, 4),
        ("· ble.zone.enter/exit", 11, INK, False, 4),
        ("· ble.geofence.breach", 11, INK, False, 4),
        ("· ble.tag.missing", 11, INK, False, 4),
        ("· ble.gw.offline", 11, INK, False, 4),
        ("· 비전 조인 stay_id", 11, INK, False),
    ])

    page += 1
    s = new_slide(prs)
    frame(s, page, "09 · 넓은 공간 BLE", "설치 배치 — 플레이그라운드 · 루프탑")
    try:
        s.shapes.add_picture(BLE_ASSETS + r"\ble_playground_gateway_plan.jpg", I(MX), I(1.75), width=I(5.9))
        s.shapes.add_picture(BLE_ASSETS + r"\ble_rooftop_gateway_plan.jpg", I(6.85), I(1.75), width=I(5.78))
    except OSError:
        text(s, MX, 2.5, CW, 0.4, [("배치도: assets/ble_tracking/", 12, GRAY, False)])
    text(s, MX, 5.85, 5.9, 0.9, [
        ("지상 플레이그라운드", 13, ACCENT, True, 3),
        ("GW 4~6대 · 실효 반경 25~35m 겹침 · 입장구 GW 필수", 11, INK, False),
    ])
    text(s, 6.85, 5.85, 5.78, 0.9, [
        ("옥상 루프탑", 13, ACCENT, True, 3),
        ("GW 3~4대 · 방수함·차양 · PoE 유선 우선", 11, INK, False),
    ])

    page += 1
    s = new_slide(prs)
    frame(s, page, "09 · 넓은 공간 BLE", "목줄 관통 케이스 · 체크인 (몸쪽 밀착)")
    try:
        s.shapes.add_picture(BLE_ASSETS + r"\ble_collar_slide_on_silicone_case.png", I(MX), I(1.75), width=I(4.2))
        s.shapes.add_picture(BLE_ASSETS + r"\ble_checkin_collar_attach_steps.jpg", I(5.15), I(1.75), width=I(7.48))
    except OSError:
        text(s, MX, 2.5, CW, 0.4, [("목줄 케이스·체크인: assets/ble_tracking/", 12, GRAY, False)])
    text(s, MX, 5.7, CW, 1.0, [
        ("1) MWB01 → 케이스 포켓  2) 목줄 띠를 좌우 루프에 관통(몸쪽 밀착)  3) 태그↔Stay 바인딩  4) 퇴장 회수·소독", 13, INK, False, 4),
        ("금지: 키링·클립 매달림형(상대견 물어뜯김). 기능: 태그 스캔 · GW 등록 · Stay.ble_tag_id.", 11, GRAY, False),
    ])

    prs.save(OUT)
    print(f"saved: {OUT}  slides={len(prs.slides)}  core={n_core}  extra={n_extra}")


if __name__ == "__main__":
    build()
